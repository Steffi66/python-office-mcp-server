"""Python bindings for the shared format-independent mutation contracts."""

import hashlib
import json
import posixpath
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pytest_bdd import given, parsers, scenarios, then, when

from office_server import OfficeServer
from tools.word_advanced_tools import _get_text_with_track_changes

SHARED = Path(__file__).parents[1] / "contracts" / "shared"
MANIFEST = json.loads((SHARED / "fixture-manifest.json").read_text())
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

scenarios("features/mutation-safety.feature")


@pytest.fixture
def ctx(tmp_path):
    return {"root": tmp_path, "server": OfficeServer()}


def snapshot(ctx):
    return {p.name: p.read_bytes() for p in ctx["root"].iterdir() if p.is_file()}


def workbook_value(path, target, data_only=False):
    wb = load_workbook(path, data_only=data_only)
    try:
        if "!" in target:
            sheet, cell = target.split("!", 1)
            return wb[sheet][cell].value
        return wb.active[target].value
    finally:
        wb.close()


@given(parsers.parse('fixture "{name}" verified against the fixture manifest'))
def fixture(ctx, name):
    record = next(f for f in MANIFEST["fixtures"] if f["id"] == name)
    data = (SHARED / record["path"]).read_bytes()
    assert hashlib.sha256(data).hexdigest() == record["sha256"]
    ctx["fixture"] = record
    ctx["source"] = ctx["root"] / name
    ctx["source"].write_bytes(data)


@given(parsers.parse('destination state is "{state}"'))
def destination(ctx, state):
    ctx["output"] = ctx["source"] if state == "source" else ctx["root"] / ("output" + ctx["source"].suffix)
    if state == "distinct-existing":
        shutil.copy2(ctx["source"], ctx["output"])


@given("source bytes, destination bytes and document-directory entries are recorded")
@given("source bytes are recorded")
def record(ctx):
    ctx["before"] = snapshot(ctx)
    ctx["entries"] = sorted(p.name for p in ctx["root"].iterdir())
    ctx["source_bytes"] = ctx["source"].read_bytes()


@given('calculation policy is "invalidate-without-recalculation"')
def no_calculation(ctx):
    ctx["calculation_policy"] = "invalidate-without-recalculation"


@when(parsers.re(r"(?P<operation>previewing this batch without committing|committing this batch with all-targets-required policy|committing this batch to the distinct destination|committing this batch with multiline wrap enabled):"))
def apply(ctx, operation, datatable, request):
    assert datatable[0] == ["target", "value_json"]
    # pytest-bdd 8 exposes raw table escapes; use the official compiler's decoded
    # table, which is the same typed input consumed by the Bun runner.
    step = request.node._office_case["steps"][request.node._office_step]
    rows = step["argument"]["dataTable"]["rows"]
    changes = [{"target": row["cells"][0]["value"], "value": json.loads(row["cells"][1]["value"])} for row in rows[1:]]
    mode = "dry_run" if operation.startswith("previewing") else "strict" if "all-targets" in operation else "safe"
    ctx["request"] = {"file_path": str(ctx["source"]), "output_path": str(ctx["output"]), "changes": changes, "mode": mode}
    ctx["result"] = ctx["server"].tool_office_patch(**ctx["request"])


@then(parsers.parse("committed change count is {count:d}"))
def committed(ctx, count):
    assert ctx["result"]["changes_applied"] == count, ctx["result"]


@then("source bytes, destination bytes and document-directory entries equal the recorded state")
def no_write(ctx):
    assert snapshot(ctx) == ctx["before"]
    assert sorted(p.name for p in ctx["root"].iterdir()) == ctx["entries"]


@then("source bytes equal the recorded state")
def unchanged_source(ctx):
    assert ctx["source"].read_bytes() == ctx["source_bytes"]


@then(parsers.re(r'the reopened (?P<which>source|destination) slide 1 (?P<shape>title|subtitle) is "(?P<text>.*)"'))
def slide_text(ctx, which, shape, text):
    slide = Presentation(ctx["source" if which == "source" else "output"]).slides[0]
    actual = slide.shapes.title.text if shape == "title" else slide.placeholders[1].text
    assert actual == text


@then(parsers.parse('the reopened source current body text is "{text}"'))
def word_text(ctx, text):
    assert "".join(_get_text_with_track_changes(p) for p in Document(ctx["source"]).paragraphs) == text


@then(parsers.parse('the reopened source active sheet cell A1 is "{text}"'))
def cell_text(ctx, text):
    assert workbook_value(ctx["source"], "A1") == text


@then(parsers.parse('the operation is refused for unmatched target "{target}"'))
def refused(ctx, target):
    assert ctx["result"]["success"] is False
    assert any(t["target"] == target for t in ctx["result"]["unmatched_targets"])


@then(parsers.parse('"{target}" is never reported applied'))
def not_applied(ctx, target):
    assert all(not item["applied"] for item in ctx["result"]["results"] if item["target"] == target)


@then(parsers.parse('the reopened destination active sheet cell A1 has value_json "{value}"'))
def multiline(ctx, value):
    assert workbook_value(ctx["output"], "A1") == json.loads('"' + value + '"')


@then("that cell resolves to a style with wrapText enabled")
def wrap(ctx):
    wb = load_workbook(ctx["output"])
    try:
        assert wb.active["A1"].alignment.wrap_text
    finally:
        wb.close()


@then("every cell style index is below the output cellXfs count")
def style_indices(ctx):
    with zipfile.ZipFile(ctx["output"]) as archive:
        count = len(ET.fromstring(archive.read("xl/styles.xml")).find(S + "cellXfs"))
        for name in archive.namelist():
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                for cell in ET.fromstring(archive.read(name)).iter(S + "c"):
                    assert 0 <= int(cell.get("s", "0")) < count


@then("all destination relationship and content-type references resolve")
def references(ctx):
    with zipfile.ZipFile(ctx["output"]) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names))
        types = ET.fromstring(archive.read("[Content_Types].xml"))
        extensions = {e.get("Extension") for e in types if e.tag.endswith("}Default")}
        overrides = {e.get("PartName", "").lstrip("/") for e in types if e.tag.endswith("}Override")}
        for name in names:
            if name != "[Content_Types].xml":
                assert name in overrides or name.rsplit(".", 1)[-1] in extensions
            if name.endswith(".rels"):
                base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
                for rel in ET.fromstring(archive.read(name)):
                    if rel.get("TargetMode") == "External":
                        continue
                    target = rel.get("Target", "").split("#", 1)[0]
                    target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    assert target in names, (name, target)


@then("destination member payloads outside the manifest change allowance are byte-identical")
def preserved_parts(ctx):
    with zipfile.ZipFile(ctx["output"]) as archive:
        assert set(archive.namelist()) == set(ctx["fixture"]["memberSha256"])
        for name, digest in ctx["fixture"]["mustPreservePayloads"].items():
            assert hashlib.sha256(archive.read(name)).hexdigest() == digest, name


@then("the reopened destination Input!A1 is numeric 10")
def numeric_input(ctx):
    value = workbook_value(ctx["output"], "Input!A1")
    assert type(value) in (int, float) and value == 10


@then(parsers.parse('the reopened destination Calc!A1 formula is "{formula}"'))
def formula(ctx, formula):
    assert workbook_value(ctx["output"], "Calc!A1") == formula


@then("the destination Calc!A1 cached value is absent or empty")
def no_cache(ctx):
    assert workbook_value(ctx["output"], "Calc!A1", data_only=True) is None


@then('calculation state is "recalculation-required"')
def calculation_state(ctx):
    assert ctx["result"]["calculation_state"] == "recalculation-required"


@then("a data-only read never returns the old cached value 2 as current")
def no_stale(ctx):
    assert workbook_value(ctx["output"], "Calc!A1", data_only=True) != 2
