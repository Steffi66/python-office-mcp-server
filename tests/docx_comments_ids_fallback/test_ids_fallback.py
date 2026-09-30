"""Run the exact commentsIds fallback case with distinct explicit and positional controls."""

from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_ids_fallback.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
W16CID = "{http://schemas.microsoft.com/office/word/2016/wordml/cid}"
EXPECTED = "0F0E0D0C"
DECOY = "ABCD1234"
ORIGINAL = "12345678"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _write_parts(path, parts):
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)


def _comments_xml(parts, comment_id):
    root = etree.fromstring(parts["word/comments.xml"])
    nodes = [node for node in root.findall(".//" + W + "comment")
             if node.get(W + "id") == comment_id]
    assert len(nodes) == 1
    paragraph = nodes[0].find(W + "p")
    assert paragraph is not None
    return root, paragraph


def _mapping_xml(comment_id, target=EXPECTED):
    root = etree.Element(W16CID + "commentsIds", nsmap={"w16cid": W16CID[1:-1]})
    for cid, para_id in (("9999", DECOY), (comment_id, target)):
        node = etree.SubElement(root, W16CID + "commentId")
        node.set(W16CID + "id", cid)
        node.set(W16CID + "paraId", para_id)
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def _read_target(reader, path, comment_id):
    response = reader.tool_word_get_comments(str(path))
    assert "error" not in response and response["comment_count"] == 1
    rows = response["comments"]
    assert [str(row["id"]) for row in rows] == [comment_id]
    assert rows[0]["text"] == "Legacy style comment"
    return rows[0]


def _prepare(path):
    document = Document()
    document.add_paragraph("Legacy mapping target")
    document.save(path)
    assert [p.text for p in Document(path).paragraphs if p.text] == ["Legacy mapping target"]
    authoring, reader = WordAdvancedTools(), WordTools()
    result = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Legacy mapping target", comment_text="Legacy style comment")
    assert result.get("success") is True
    with_id = reader.tool_word_get_comments(str(path))
    assert "error" not in with_id and with_id["comment_count"] == 1
    comment_id = str(with_id["comments"][0]["id"])
    assert comment_id != "9999"
    assert with_id["comments"][0]["text"] == "Legacy style comment"
    parts = _parts(path)
    xml, paragraph = _comments_xml(parts, comment_id)
    assert paragraph.get(W14 + "paraId") is None
    original_para_id = ORIGINAL
    paragraph.set(W14 + "paraId", original_para_id)
    parts["word/comments.xml"] = etree.tostring(xml, encoding="UTF-8", xml_declaration=True)
    _write_parts(path, parts)
    assert _read_target(reader, path, comment_id)["para_id"] == original_para_id
    paragraph.attrib.pop(W14 + "paraId")
    assert paragraph.get(W14 + "paraId") is None
    parts["word/comments.xml"] = etree.tostring(xml, encoding="UTF-8", xml_declaration=True, standalone=True)
    return reader, comment_id, parts, original_para_id


def test_canonical_comments_ids_fallback(ids_fallback_case, tmp_path, request):
    case = ids_fallback_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_ids_fallback_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "ids-fallback.docx"
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                document = Document()
                document.add_paragraph("Legacy mapping target")
                document.save(path)
                assert [p.text for p in Document(path).paragraphs if p.text] == ["Legacy mapping target"]
                authoring, reader = WordAdvancedTools(), WordTools()
                result = authoring.tool_word_add_comment(str(path), "Legacy mapping target", "Legacy style comment")
                assert result.get("success") is True
                initial = reader.tool_word_get_comments(str(path))
                assert "error" not in initial and initial["comment_count"] == 1
                comment_id = str(initial["comments"][0]["id"])
                assert comment_id != "9999"
                assert initial["comments"][0]["text"] == "Legacy style comment"
                parts = _parts(path)
                xml, paragraph = _comments_xml(parts, comment_id)
                assert paragraph.get(W14 + "paraId") is None
                original_para_id = ORIGINAL
                paragraph.set(W14 + "paraId", original_para_id)
                parts["word/comments.xml"] = etree.tostring(xml, encoding="UTF-8", xml_declaration=True)
                _write_parts(path, parts)
                assert _read_target(reader, path, comment_id)["para_id"] == original_para_id
                case["commentId"] = comment_id
            elif index == 1:
                paragraph.attrib.pop(W14 + "paraId")
                assert paragraph.get(W14 + "paraId") is None
                parts["word/comments.xml"] = etree.tostring(
                    xml, encoding="UTF-8", xml_declaration=True, standalone=True)
            elif index == 2:
                assert "word/commentsIds.xml" not in parts
                parts["word/commentsIds.xml"] = _mapping_xml(comment_id)
                mapping = etree.fromstring(parts["word/commentsIds.xml"])
                nodes = mapping.findall(W16CID + "commentId")
                assert [(node.get(W16CID + "id"), node.get(W16CID + "paraId")) for node in nodes] == [
                    ("9999", DECOY), (comment_id, EXPECTED)]
                _write_parts(path, parts)
                saved = _parts(path)
                assert saved["word/commentsIds.xml"] == parts["word/commentsIds.xml"]
                _, no_id = _comments_xml(saved, comment_id)
                assert no_id.get(W14 + "paraId") is None
            elif index == 3:
                before = path.read_bytes()
                returned = _read_target(reader, path, comment_id)
                assert path.read_bytes() == before
            else:
                assert returned["para_id"] == EXPECTED
                assert returned["para_id"] != original_para_id and returned["para_id"] != DECOY
                case["resolvedParaId"] = returned["para_id"]
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_comment_id_mapping_not_first_positional_entry(tmp_path):
    path = tmp_path / "distinct-mapping.docx"
    reader, comment_id, parts, original = _prepare(path)
    parts["word/commentsIds.xml"] = _mapping_xml(comment_id)
    _write_parts(path, parts)
    assert _read_target(reader, path, comment_id)["para_id"] == EXPECTED
    # Removing only the matching ID leaves the positional decoy in place.
    xml = etree.fromstring(parts["word/commentsIds.xml"])
    xml.remove(xml[-1])
    parts["word/commentsIds.xml"] = etree.tostring(xml, encoding="UTF-8", xml_declaration=True)
    _write_parts(path, parts)
    assert _read_target(reader, path, comment_id)["para_id"] == DECOY
    assert DECOY != original


def test_original_para_id_and_missing_map_controls(tmp_path):
    path = tmp_path / "without-map.docx"
    reader, comment_id, parts, original = _prepare(path)
    _write_parts(path, parts)
    assert _read_target(reader, path, comment_id)["para_id"] is None
    xml, paragraph = _comments_xml(parts, comment_id)
    paragraph.set(W14 + "paraId", original)
    parts["word/comments.xml"] = etree.tostring(xml, encoding="UTF-8", xml_declaration=True)
    parts["word/commentsIds.xml"] = _mapping_xml(comment_id)
    _write_parts(path, parts)
    assert _read_target(reader, path, comment_id)["para_id"] == original
    assert original != EXPECTED
