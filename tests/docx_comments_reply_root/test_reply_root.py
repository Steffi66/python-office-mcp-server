"""Execute the exact authored reply-to-root case with separate controls."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_reply_root.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _comments(reader, path, count):
    response = reader.tool_word_get_comments(str(path), filter="all")
    assert "error" not in response and response.get("filter") == "all"
    rows = response["comments"]
    assert response["comment_count"] == len(rows) == count
    return rows


def _create_root(path, authoring, reader):
    doc = Document()
    doc.add_paragraph("Resolve this thread")
    doc.save(path)
    assert path.is_file() and [p.text for p in Document(path).paragraphs if p.text] == ["Resolve this thread"]
    response = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Resolve this thread", comment_text="Root comment")
    assert response.get("success") is True
    roots = _comments(reader, path, 1)
    assert roots[0]["text"] == "Root comment" and roots[0]["done"] is False
    root_id = str(roots[0]["id"])
    document = etree.fromstring(_parts(path)["word/document.xml"])
    targets = [node for node in document.findall(".//" + W + "body/" + W + "p")
               if "".join(node.itertext()) == "Resolve this thread"]
    assert len(targets) == 1
    assert [n.get(W + "id") for n in targets[0].findall(".//" + W + "commentRangeStart")] == [root_id]
    assert [n.get(W + "id") for n in targets[0].findall(".//" + W + "commentRangeEnd")] == [root_id]
    return root_id


def _add_reply(path, reader, root_id):
    result = reader.tool_word_reply_to_comment(str(path), root_id, "Reply comment")
    assert result.get("success") is True
    reply_id = str(result["reply_comment_id"])
    assert reply_id != root_id
    rows = _comments(reader, path, 2)
    assert [(str(r["id"]), r["text"], r["done"]) for r in rows] == [
        (root_id, "Root comment", False), (reply_id, "Reply comment", False)]
    assert rows[1]["is_reply"] is True and str(rows[1]["parent_id"]) == root_id
    comments = etree.fromstring(_parts(path)["word/comments.xml"])
    root, reply = comments.findall(".//" + W + "comment")
    assert root.get(W + "id") == root_id and reply.get(W + "id") == reply_id
    root_para = root.find(W + "p")
    reply_para = reply.find(W + "p")
    assert root_para is not None and reply_para is not None
    assert root_para.get(W14 + "paraId") and reply_para.get(W14 + "paraId")
    assert reply_para.get(W14 + "paraIdParent") == root_para.get(W14 + "paraId")
    return result, reply_id


def test_canonical_reply_resolves_root(reply_root_case, tmp_path, request):
    case = reply_root_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_reply_root_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "reply-root.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                doc = Document()
                doc.add_paragraph("Resolve this thread")
                doc.save(path)
                assert path.is_file() and [p.text for p in Document(path).paragraphs if p.text] == ["Resolve this thread"]
            elif index == 1:
                result = authoring.tool_word_add_comment(
                    file_path=str(path), target_text="Resolve this thread", comment_text="Root comment")
                assert result.get("success") is True
                roots = _comments(reader, path, 1)
                assert roots[0]["text"] == "Root comment" and roots[0]["done"] is False
                root_id = str(roots[0]["id"])
                body = etree.fromstring(_parts(path)["word/document.xml"])
                targets = [p for p in body.findall(".//" + W + "body/" + W + "p")
                           if "".join(p.itertext()) == "Resolve this thread"]
                assert len(targets) == 1
                assert [n.get(W + "id") for n in targets[0].findall(".//" + W + "commentRangeStart")] == [root_id]
                assert [n.get(W + "id") for n in targets[0].findall(".//" + W + "commentRangeEnd")] == [root_id]
                case["rootId"] = root_id
            elif index == 2:
                reply_response, reply_id = _add_reply(path, reader, root_id)
                case["replyId"] = reply_id
            elif index == 3:
                resolve_response = reader.tool_word_resolve_comment(str(path), reply_id, True)
                assert path.is_file()
            elif index == 4:
                assert reply_response.get("success") is True and resolve_response.get("success") is True
                assert resolve_response.get("done") is True
            elif index == 5:
                assert str(resolve_response.get("thread_root_comment_id")) == root_id
                assert str(resolve_response.get("comment_id")) == reply_id
            else:
                fresh = _comments(reader, path, 2)
                assert [(str(r["id"]), r["text"], r["done"]) for r in fresh] == [
                    (root_id, "Root comment", True), (reply_id, "Reply comment", False)]
                assert fresh[1]["is_reply"] is True and str(fresh[1]["parent_id"]) == root_id
                changed = _parts(path)
                extension = etree.fromstring(changed["word/commentsExtended.xml"])
                assert any(node.get(W15 + "paraId") == str(fresh[0]["para_id"])
                           and node.get(W15 + "done") == "1" for node in extension.findall(W15 + "commentEx"))
                case["rootDone"] = True
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_root_target_is_distinct_from_reply_target(tmp_path):
    path = tmp_path / "direct-root.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    root_id = _create_root(path, authoring, reader)
    _, reply_id = _add_reply(path, reader, root_id)
    response = reader.tool_word_resolve_comment(str(path), root_id, True)
    assert response.get("success") is True
    assert str(response.get("comment_id")) == root_id and str(response.get("thread_root_comment_id")) == root_id
    assert [(str(r["id"]), r["done"]) for r in _comments(reader, path, 2)] == [
        (root_id, True), (reply_id, False)]
    prior = path.read_bytes()
    again = reader.tool_word_resolve_comment(str(path), root_id, True)
    assert again.get("success") is True and again.get("unchanged") is True
    assert path.read_bytes() == prior


def test_unknown_reply_id_does_not_publish(tmp_path):
    path = tmp_path / "missing.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    root_id = _create_root(path, authoring, reader)
    _, reply_id = _add_reply(path, reader, root_id)
    source = path.read_bytes()
    unknown = str(max(int(root_id), int(reply_id)) + 100)
    response = reader.tool_word_resolve_comment(str(path), unknown, True)
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source
    assert [(str(r["id"]), r["done"]) for r in _comments(reader, path, 2)] == [
        (root_id, False), (reply_id, False)]
