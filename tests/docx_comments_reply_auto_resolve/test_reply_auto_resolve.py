"""Run exact authored reply auto-resolution with separate root-selection controls."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_reply_auto_resolve.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _rows(reader, path, selected, count):
    response = reader.tool_word_get_comments(str(path), filter=selected)
    assert "error" not in response and response.get("filter") == selected
    assert response["comment_count"] == len(response["comments"]) == count
    return response["comments"]


def _root(path, authoring, reader):
    doc = Document()
    doc.add_paragraph("Auto resolve target")
    doc.save(path)
    assert [p.text for p in Document(path).paragraphs if p.text] == ["Auto resolve target"]
    response = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Auto resolve target", comment_text="Needs action")
    assert response.get("success") is True
    roots = _rows(reader, path, "all", 1)
    assert roots[0]["text"] == "Needs action" and roots[0]["done"] is False
    root_id = str(roots[0]["id"])
    body = etree.fromstring(_parts(path)["word/document.xml"])
    paragraphs = [p for p in body.findall(".//" + W + "body/" + W + "p")
                  if "".join(p.itertext()) == "Auto resolve target"]
    assert len(paragraphs) == 1
    assert [node.get(W + "id") for node in paragraphs[0].findall(".//" + W + "commentRangeStart")] == [root_id]
    assert [node.get(W + "id") for node in paragraphs[0].findall(".//" + W + "commentRangeEnd")] == [root_id]
    return root_id


def _reply_relation(path, reader, root_id, reply_id):
    assert reply_id != root_id
    flat = _rows(reader, path, "all", 2)
    assert [(str(row["id"]), row["text"], row["done"]) for row in flat] == [
        (root_id, "Needs action", True), (reply_id, "Done now", False)]
    assert flat[1]["is_reply"] is True and str(flat[1]["parent_id"]) == root_id
    comments = etree.fromstring(_parts(path)["word/comments.xml"])
    root, reply = comments.findall(".//" + W + "comment")
    assert [root.get(W + "id"), reply.get(W + "id")] == [root_id, reply_id]
    root_para, reply_para = root.find(W + "p"), reply.find(W + "p")
    assert root_para is not None and reply_para is not None
    assert root_para.get(W14 + "paraId") and reply_para.get(W14 + "paraId")
    assert reply_para.get(W14 + "paraIdParent") == root_para.get(W14 + "paraId")
    return flat[0]["para_id"]


def test_canonical_reply_auto_resolve(reply_auto_resolve_case, tmp_path, request):
    case = reply_auto_resolve_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_reply_auto_resolve_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "auto-resolve.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                root_id = _root(path, authoring, reader)
                case["rootId"] = root_id
            elif index == 1:
                response = reader.tool_word_reply_to_comment(
                    file_path=str(path), comment_id=root_id, text="Done now", auto_resolve=True)
                assert path.is_file()
                reply_id = str(response["reply_comment_id"])
                assert reply_id != root_id
                case["replyId"] = reply_id
            elif index == 2:
                assert response.get("success") is True and response.get("resolved") is True
                assert response.get("auto_resolve") is True
                assert str(response.get("parent_comment_id")) == root_id
                assert str(response.get("thread_root_comment_id")) == root_id
            else:
                resolved = _rows(reader, path, "resolved", 1)
                assert [(str(row["id"]), row["text"], row["done"]) for row in resolved] == [
                    (root_id, "Needs action", True)]
                para_id = _reply_relation(path, reader, root_id, reply_id)
                extension = etree.fromstring(_parts(path)["word/commentsExtended.xml"])
                assert any(node.get(W15 + "paraId") == str(para_id) and node.get(W15 + "done") == "1"
                           for node in extension.findall(W15 + "commentEx"))
                assert _rows(reader, path, "open", 1)[0]["id"] == reply_id
                case["resolvedRootId"] = root_id
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_reply_without_auto_resolve_keeps_root_open(tmp_path):
    path = tmp_path / "no-auto-resolve.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    root_id = _root(path, authoring, reader)
    response = reader.tool_word_reply_to_comment(str(path), root_id, "Done now", auto_resolve=False)
    assert response.get("success") is True and response.get("auto_resolve") is False
    assert response.get("resolved") is not True
    reply_id = str(response["reply_comment_id"])
    rows = _rows(reader, path, "all", 2)
    assert [(str(row["id"]), row["done"]) for row in rows] == [(root_id, False), (reply_id, False)]
    assert _rows(reader, path, "resolved", 0) == []
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "open", 2)] == [
        (root_id, False), (reply_id, False)]


def test_second_thread_is_only_resolved_target(tmp_path):
    path = tmp_path / "second-thread.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    doc = Document()
    doc.add_paragraph("First target")
    doc.add_paragraph("Second target")
    doc.save(path)
    for target, text in (("First target", "First root"), ("Second target", "Second root")):
        result = authoring.tool_word_add_comment(str(path), target, text)
        assert result.get("success") is True
    roots = _rows(reader, path, "all", 2)
    first, second = (str(row["id"]) for row in roots)
    assert first != second and [row["text"] for row in roots] == ["First root", "Second root"]
    response = reader.tool_word_reply_to_comment(str(path), second, "Second reply", auto_resolve=True)
    assert response.get("success") is True and response.get("resolved") is True
    assert str(response.get("parent_comment_id")) == second
    assert str(response.get("thread_root_comment_id")) == second
    reply_id = str(response["reply_comment_id"])
    assert reply_id not in (first, second)
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "resolved", 1)] == [(second, True)]
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "open", 2)] == [
        (first, False), (reply_id, False)]
    all_rows = _rows(reader, path, "all", 3)
    assert str(all_rows[2]["parent_id"]) == second and all_rows[2]["is_reply"] is True


def test_unknown_target_does_not_publish(tmp_path):
    path = tmp_path / "unknown.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    root_id = _root(path, authoring, reader)
    source = path.read_bytes()
    response = reader.tool_word_reply_to_comment(str(path), str(int(root_id) + 1000), "Done now", auto_resolve=True)
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source
    assert _rows(reader, path, "resolved", 0) == []
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "all", 1)] == [(root_id, False)]
