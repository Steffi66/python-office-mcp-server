"""Execute the exact authored two-comment case; controls are not sibling credit."""

from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from tests.docx_comments_mixed_done.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"


def _archive_parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _comments(response):
    assert "error" not in response and response.get("filter") == "all"
    records = response["comments"]
    assert response["comment_count"] == len(records) == 2
    assert len({str(record["id"]) for record in records}) == 2
    return records


def _authored_document(path, authoring, reader):
    doc = Document()
    doc.add_paragraph("Alpha target")
    doc.add_paragraph("Beta target")
    doc.save(path)
    assert path.is_file() and [p.text for p in Document(path).paragraphs if p.text] == ["Alpha target", "Beta target"]
    alpha = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Alpha target", comment_text="Comment alpha", author="Manuel")
    beta = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Beta target", comment_text="Comment beta", author="Rui Carmo")
    assert alpha.get("success") is True and beta.get("success") is True
    rows = _comments(reader.tool_word_get_comments(str(path), filter="all"))
    assert [(r["text"], r["author"], r["done"]) for r in rows] == [
        ("Comment alpha", "Manuel", False), ("Comment beta", "Rui Carmo", False)]
    return rows


def test_canonical_authored_mixed_done(mixed_done_case, tmp_path, request):
    case = mixed_done_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_mixed_done_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "mixed.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                doc = Document()
                doc.add_paragraph("Alpha target")
                doc.add_paragraph("Beta target")
                doc.save(path)
                assert [p.text for p in Document(path).paragraphs if p.text] == ["Alpha target", "Beta target"]
            elif index == 1:
                alpha = authoring.tool_word_add_comment(
                    file_path=str(path), target_text="Alpha target", comment_text="Comment alpha", author="Manuel")
                beta = authoring.tool_word_add_comment(
                    file_path=str(path), target_text="Beta target", comment_text="Comment beta", author="Rui Carmo")
                assert alpha.get("success") is True and beta.get("success") is True
                assert path.is_file()
            elif index == 2:
                before = _comments(reader.tool_word_get_comments(str(path), filter="all"))
                assert [(r["text"], r["author"], r["done"]) for r in before] == [
                    ("Comment alpha", "Manuel", False), ("Comment beta", "Rui Carmo", False)]
                first_id, second_id = (str(r["id"]) for r in before)
                assert first_id != second_id
                original_parts = _archive_parts(path)
                assert "word/comments.xml" in original_parts
                body = etree.fromstring(original_parts["word/document.xml"])
                w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                paragraphs = body.findall(".//" + w + "body/" + w + "p")
                targets = [p for p in paragraphs if "".join(p.itertext()) in ("Alpha target", "Beta target")]
                assert ["".join(p.itertext()) for p in targets] == ["Alpha target", "Beta target"]
                for target, expected_id in zip(targets, (first_id, second_id)):
                    starts = target.findall(".//" + w + "commentRangeStart")
                    ends = target.findall(".//" + w + "commentRangeEnd")
                    assert [n.get(w + "id") for n in starts] == [expected_id]
                    assert [n.get(w + "id") for n in ends] == [expected_id]
                case["returnedOrder"] = [first_id, second_id]
            elif index == 3:
                resolved = reader.tool_word_resolve_comment(str(path), first_id, True)
                assert path.is_file()
            elif index == 4:
                assert resolved.get("success") is True and resolved.get("done") is True
                assert str(resolved.get("thread_root_comment_id")) == first_id
            else:
                observed = _comments(reader.tool_word_get_comments(str(path), filter="all"))
                assert [(str(r["id"]), r["text"], r["author"], r["done"]) for r in observed] == [
                    (first_id, "Comment alpha", "Manuel", True),
                    (second_id, "Comment beta", "Rui Carmo", False)]
                changed_parts = _archive_parts(path)
                # Authored comments may acquire a missing paraId on resolution;
                # this case promises mixed state, not lexical member custody.
                old_comments = etree.fromstring(original_parts["word/comments.xml"])
                new_comments = etree.fromstring(changed_parts["word/comments.xml"])
                assert [(n.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"),
                         "".join(n.itertext())) for n in old_comments] == [
                    (n.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"),
                     "".join(n.itertext())) for n in new_comments]
                extension = etree.fromstring(changed_parts["word/commentsExtended.xml"])
                states = {str(node.get(W15 + "paraId")): node.get(W15 + "done")
                          for node in extension.findall(W15 + "commentEx")}
                assert states[str(observed[0]["para_id"])] == "1"
                assert str(observed[1]["para_id"]) not in states or states[str(observed[1]["para_id"])] == "0"
                case["observedMixedState"] = {first_id: True, second_id: False}
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_opposite_id_and_noop_controls(tmp_path):
    path = tmp_path / "controls.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    original = _authored_document(path, authoring, reader)
    first_id, second_id = (str(r["id"]) for r in original)
    result = reader.tool_word_resolve_comment(str(path), second_id, True)
    assert result.get("success") is True and str(result.get("thread_root_comment_id")) == second_id
    mixed = _comments(reader.tool_word_get_comments(str(path), filter="all"))
    assert [r["done"] for r in mixed] == [False, True]
    prior_archive = path.read_bytes()
    again = reader.tool_word_resolve_comment(str(path), second_id, True)
    assert again.get("success") is True and again.get("unchanged") is True
    assert path.read_bytes() == prior_archive
    assert [r["done"] for r in _comments(reader.tool_word_get_comments(str(path), filter="all"))] == [False, True]
    assert first_id != second_id


def test_missing_id_does_not_publish(tmp_path):
    path = tmp_path / "missing.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    rows = _authored_document(path, authoring, reader)
    source = path.read_bytes()
    unknown = str(max(int(row["id"]) for row in rows) + 100)
    result = reader.tool_word_resolve_comment(str(path), unknown, True)
    assert "error" in result and result.get("success") is not True
    assert path.read_bytes() == source
