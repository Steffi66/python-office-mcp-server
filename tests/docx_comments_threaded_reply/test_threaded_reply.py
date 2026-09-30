"""Run the exact authored threaded-reply read with separate wrong-group controls."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_threaded_reply.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _flat(reader, path, count):
    response = reader.tool_word_get_comments(str(path), format="flat")
    assert "error" not in response and response.get("filter") == "all"
    assert response["comment_count"] == len(response["comments"]) == count
    return response["comments"]


def _threaded(reader, path, roots, entries):
    before = path.read_bytes()
    response = reader.tool_word_get_comments(str(path), format="threaded")
    assert path.read_bytes() == before and "error" not in response
    assert response.get("filter") == "all"
    assert response["comment_count"] == len(response["comments"]) == entries
    assert response["flat"] == response["comments"]
    assert response["thread_count"] == len(response["threads"]) == roots
    return response


def _create(path, authoring, reader):
    doc = Document()
    doc.add_paragraph("Thread me")
    doc.save(path)
    assert [p.text for p in Document(path).paragraphs if p.text] == ["Thread me"]
    created = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Thread me", comment_text="Root")
    assert created.get("success") is True
    rows = _flat(reader, path, 1)
    assert rows[0]["text"] == "Root" and rows[0]["done"] is False
    root_id = str(rows[0]["id"])
    body = etree.fromstring(_parts(path)["word/document.xml"])
    paragraphs = [p for p in body.findall(".//" + W + "body/" + W + "p")
                  if "".join(p.itertext()) == "Thread me"]
    assert len(paragraphs) == 1
    assert [node.get(W + "id") for node in paragraphs[0].findall(".//" + W + "commentRangeStart")] == [root_id]
    assert [node.get(W + "id") for node in paragraphs[0].findall(".//" + W + "commentRangeEnd")] == [root_id]
    return root_id


def _reply(path, reader, root_id):
    response = reader.tool_word_reply_to_comment(str(path), root_id, "Reply")
    assert response.get("success") is True
    reply_id = str(response["reply_comment_id"])
    assert reply_id != root_id and str(response["parent_comment_id"]) == root_id
    rows = _flat(reader, path, 2)
    assert [(str(row["id"]), row["text"], row["done"]) for row in rows] == [
        (root_id, "Root", False), (reply_id, "Reply", False)]
    assert rows[1]["is_reply"] is True and str(rows[1]["parent_id"]) == root_id
    comments = etree.fromstring(_parts(path)["word/comments.xml"])
    roots = comments.findall(".//" + W + "comment")
    assert [node.get(W + "id") for node in roots] == [root_id, reply_id]
    root_para, reply_para = (node.find(W + "p") for node in roots)
    assert root_para is not None and reply_para is not None
    assert root_para.get(W14 + "paraId") and reply_para.get(W14 + "paraId")
    assert reply_para.get(W14 + "paraIdParent") == root_para.get(W14 + "paraId")
    return response, reply_id


def test_canonical_threaded_reply(threaded_reply_case, tmp_path, request):
    case = threaded_reply_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_threaded_reply_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "threaded-reply.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                root_id = _create(path, authoring, reader)
                case["rootId"] = root_id
            elif index == 1:
                reply_response, reply_id = _reply(path, reader, root_id)
                case["replyId"] = reply_id
            elif index == 2:
                threaded = _threaded(reader, path, 1, 2)
            elif index == 3:
                assert reply_response.get("success") is True
                assert reply_response.get("auto_resolve") is False
            elif index == 4:
                assert "threads" in threaded and threaded["thread_count"] >= 1
                assert threaded["thread_count"] == 1
            elif index == 5:
                thread = threaded["threads"][0]
                assert str(thread["root"]["id"]) == root_id
                assert thread["root"]["text"] == "Root" and thread["root"]["done"] is False
            else:
                assert [(str(row["id"]), row["text"], row["done"]) for row in thread["replies"]] == [
                    (reply_id, "Reply", False)]
                assert thread["replies"][0]["is_reply"] is True
                assert str(thread["replies"][0]["parent_id"]) == root_id
                case["threadGrouped"] = True
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_two_roots_keep_reply_in_second_thread(tmp_path):
    path = tmp_path / "two-threads.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    doc = Document()
    doc.add_paragraph("First target")
    doc.add_paragraph("Second target")
    doc.save(path)
    for target, text in (("First target", "First root"), ("Second target", "Second root")):
        created = authoring.tool_word_add_comment(str(path), target, text)
        assert created.get("success") is True
    roots = _flat(reader, path, 2)
    first_id, second_id = (str(row["id"]) for row in roots)
    assert first_id != second_id and [row["text"] for row in roots] == ["First root", "Second root"]
    reply = reader.tool_word_reply_to_comment(str(path), second_id, "Second reply")
    assert reply.get("success") is True and str(reply.get("parent_comment_id")) == second_id
    reply_id = str(reply["reply_comment_id"])
    assert reply_id not in (first_id, second_id)
    output = _threaded(reader, path, 2, 3)
    first, second = output["threads"]
    assert str(first["root"]["id"]) == first_id and first["replies"] == []
    assert str(second["root"]["id"]) == second_id
    assert [(str(row["id"]), row["text"]) for row in second["replies"]] == [(reply_id, "Second reply")]
    assert str(second["replies"][0]["parent_id"]) == second_id
    assert [(str(row["id"]), row["text"]) for row in output["flat"]] == [
        (first_id, "First root"), (second_id, "Second root"), (reply_id, "Second reply")]


def test_unknown_reply_id_does_not_publish(tmp_path):
    path = tmp_path / "unknown.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    root_id = _create(path, authoring, reader)
    source = path.read_bytes()
    response = reader.tool_word_reply_to_comment(str(path), str(int(root_id) + 1000), "Unknown reply")
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source
    assert _threaded(reader, path, 1, 1)["threads"][0]["replies"] == []
