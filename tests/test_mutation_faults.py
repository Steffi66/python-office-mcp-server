"""Failure injection and bounded concurrency tests for the publication boundary."""

import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tools import mutation


def setup_files(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 0
    workbook.save(source)
    destination = tmp_path / "output.xlsx"
    destination.write_bytes(b"previous output")
    return source, destination


def apply_increment(staged):
    workbook = load_workbook(staged)
    workbook.active["A1"].value += 1
    workbook.save(staged)
    workbook.close()
    return {"success": True, "results": [{"target": "A1", "success": True}]}


@pytest.mark.parametrize("failure", ["apply", "validation", "replace"])
def test_failed_phase_preserves_existing_destination_and_cleans(tmp_path, monkeypatch, failure):
    source, output = setup_files(tmp_path)
    before = {p.name: p.read_bytes() for p in tmp_path.iterdir()}

    def fail(*_args):
        raise OSError("injected failure")

    if failure == "validation":
        monkeypatch.setattr(mutation, "validate_staged_document", fail)
    if failure == "replace":
        monkeypatch.setattr(mutation.os, "replace", fail)
    result = mutation.stage_patch(str(source), str(output), "strict", 1, fail if failure == "apply" else apply_increment)
    assert result["success"] is False
    assert "injected failure" in result["error"]
    assert result["changes_applied"] == 0
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == before
    assert not mutation._path_locks


@pytest.mark.parametrize("changed", ["source", "destination"])
def test_external_change_is_not_overwritten(tmp_path, changed):
    source, output = setup_files(tmp_path)
    untouched = source if changed == "destination" else output
    original = untouched.read_bytes()

    def apply(staged):
        result = apply_increment(staged)
        (source if changed == "source" else output).write_bytes(b"external edit")
        return result

    result = mutation.stage_patch(str(source), str(output), "strict", 1, apply)
    assert result["success"] is False
    assert "changed before commit" in result["error"]
    assert (source if changed == "source" else output).read_bytes() == b"external edit"
    assert untouched.read_bytes() == original
    assert not any(p.is_dir() for p in tmp_path.iterdir())


def test_same_source_writers_are_serialised_without_lost_updates(tmp_path):
    source, _ = setup_files(tmp_path)
    ready = threading.Event()
    release = threading.Event()
    second_entered = threading.Event()

    def first(staged):
        ready.set()
        assert release.wait(5)
        return apply_increment(staged)

    def second(staged):
        second_entered.set()
        return apply_increment(staged)

    with ThreadPoolExecutor(max_workers=2) as pool:
        a = pool.submit(mutation.stage_patch, str(source), None, "strict", 1, first)
        assert ready.wait(5)
        b = pool.submit(mutation.stage_patch, str(source), None, "strict", 1, second)
        try:
            assert not second_entered.wait(0.1)
        finally:
            release.set()
        assert a.result(timeout=5)["changes_applied"] == 1
        assert b.result(timeout=5)["changes_applied"] == 1
    workbook = load_workbook(source)
    try:
        assert workbook.active["A1"].value == 2
    finally:
        workbook.close()
    assert not mutation._path_locks


def test_symlink_source_mutates_target_without_replacing_link(tmp_path):
    source, _ = setup_files(tmp_path)
    alias = tmp_path / "alias.xlsx"
    alias.symlink_to(source)
    result = mutation.stage_patch(str(alias), None, "strict", 1, apply_increment)
    assert result["changes_applied"] == 1, result
    assert alias.is_symlink()
    workbook = load_workbook(source)
    try:
        assert workbook.active["A1"].value == 1
    finally:
        workbook.close()


def test_hard_links_are_refused_without_mutation(tmp_path):
    source, _ = setup_files(tmp_path)
    alias = tmp_path / "alias.xlsx"
    alias.hardlink_to(source)
    before = source.read_bytes()
    result = mutation.stage_patch(str(alias), None, "strict", 1, apply_increment)
    assert result["changes_applied"] == 0
    assert "Hard-linked" in result["error"]
    assert source.read_bytes() == before


def test_invalid_staged_archive_never_replaces_output(tmp_path):
    source, output = setup_files(tmp_path)

    def corrupt(staged):
        Path(staged).write_bytes(b"broken zip")
        return {"success": True, "results": [{"target": "A1", "success": True}]}

    result = mutation.stage_patch(str(source), str(output), "strict", 1, corrupt)
    assert result["changes_applied"] == 0
    assert output.read_bytes() == b"previous output"
