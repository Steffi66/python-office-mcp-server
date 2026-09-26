"""Optional independent calculation/rendering lane; required-oracle CI must provision LO."""

import json
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from office_server import OfficeServer

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
pytestmark = [pytest.mark.oracle, pytest.mark.skipif(not SOFFICE, reason="LibreOffice unavailable: calculation/rendering unverified")]
ROOT = Path(__file__).parents[1]


def convert(tmp_path, source, fmt):
    destination = tmp_path / "converted"
    destination.mkdir(exist_ok=True)
    profile = (tmp_path / "lo-profile").as_uri()
    result = subprocess.run([SOFFICE, "-env:UserInstallation=" + profile, "--headless", "--convert-to", fmt,
                             "--outdir", str(destination), str(source)], capture_output=True, text=True, timeout=90)
    evidence = ROOT / "test-results" / "oracle"
    evidence.mkdir(parents=True, exist_ok=True)
    (evidence / (source.stem + ".json")).write_text(json.dumps({"returncode": result.returncode, "stdout": result.stdout, "stderr": result.stderr}))
    assert result.returncode == 0, result.stderr
    output = destination / (source.stem + "." + fmt)
    assert output.exists(), result.stdout + result.stderr
    return output


def test_libreoffice_recalculates_invalidated_cross_sheet_cache(tmp_path):
    source = tmp_path / "calculation.xlsx"
    shutil.copy2(ROOT / "tests/contracts/shared/fixtures/cross-sheet-cache.xlsx", source)
    result = OfficeServer().tool_office_patch(str(source), [{"target": "Input!A1", "value": 10}])
    assert result["calculation_state"] == "recalculation-required"
    output = convert(tmp_path, source, "xlsx")
    workbook = load_workbook(output, data_only=True)
    try:
        assert workbook["Calc"]["A1"].value == 20
    finally:
        workbook.close()


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_libreoffice_renders_edited_document_to_pdf(tmp_path, suffix):
    source = tmp_path / ("render" + suffix)
    if suffix == ".docx":
        document = Document()
        document.add_paragraph("<Present>")
        target = "<Present>"
    else:
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = "Original"
        target = "slide:1/title"
    document.save(source)
    result = OfficeServer().tool_office_patch(str(source), [{"target": target, "value": "Verified render"}])
    assert result["changes_applied"] == 1
    output = convert(tmp_path, source, "pdf")
    assert output.read_bytes().startswith(b"%PDF-")
    assert output.stat().st_size > 1000
