"""Actual loopback Streamable HTTP sessions, bounded credentials and Office edits."""

import asyncio
import http.client
import json
import os
import socket
import subprocess
import sys
import time
from contextlib import contextmanager
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from office_server import OfficeServer

VERSION = "2025-03-26"
TOKEN = "synthetic-office-test-token"
ROOT = Path(__file__).parents[1]


@contextmanager
def running(tmp_path, *flags, token=TOKEN):
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    python = os.environ.get("OFFICE_MCP_TEST_PYTHON", sys.executable)
    command = [python, "-m", "office_server"] if "OFFICE_MCP_TEST_PYTHON" in os.environ else [python, str(ROOT / "office_server.py")]
    env = os.environ.copy()
    env.pop("OFFICE_MCP_HTTP_TOKEN", None)
    if token:
        env["OFFICE_MCP_HTTP_TOKEN"] = token
    with (tmp_path / "server.log").open("w+") as log:
        proc = subprocess.Popen(command + ["--port", str(port), *flags], cwd=tmp_path, env=env, stdout=log, stderr=log)
        try:
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                if proc.poll() is not None:
                    log.seek(0)
                    raise AssertionError(log.read())
                try:
                    with socket.create_connection(("127.0.0.1", port), timeout=.1):
                        break
                except OSError:
                    time.sleep(.02)
            else:
                raise AssertionError("server startup timed out")
            yield port
        finally:
            proc.terminate()
            try:
                proc.wait(timeout=3)
            except subprocess.TimeoutExpired:
                proc.kill()
                proc.wait(timeout=3)


@pytest.fixture
def server(tmp_path):
    with running(tmp_path, "--http", "--max-request-bytes", "65536") as port:
        yield port


def rpc(method, params=None, ident=1):
    return {"jsonrpc": "2.0", "id": ident, "method": method, "params": params or {}}


def post(connection, message, session=None, token=TOKEN, version=VERSION, **extra):
    headers = {"Content-Type": "application/json", "Accept": "application/json"}
    if token:
        headers["Authorization"] = "Bearer " + token
    if message.get("method") != "initialize" and version:
        headers["MCP-Protocol-Version"] = version
    if session:
        headers["Mcp-Session-Id"] = session
    headers.update(extra)
    connection.request("POST", "/mcp", json.dumps(message), headers)
    response = connection.getresponse()
    body = response.read()
    return response, json.loads(body) if body else None


def initialize(connection):
    response, body = post(connection, rpc("initialize", {"protocolVersion": VERSION, "capabilities": {}, "clientInfo": {"name": "office-tests", "version": "1"}}))
    assert response.status == 200, body
    assert body["result"]["protocolVersion"] == VERSION
    session = response.getheader("Mcp-Session-Id")
    assert session
    return session


def test_persistent_session_office_mutation_delete_and_errors(server, tmp_path):
    conn = http.client.HTTPConnection("127.0.0.1", server, timeout=5)
    try:
        session = initialize(conn)
        sock = conn.sock
        response, tools = post(conn, rpc("tools/list"), session)
        assert response.status == 200 and conn.sock is sock
        assert "office_patch" in [t["name"] for t in tools["result"]["tools"]]
        source, output = tmp_path / "source.xlsx", tmp_path / "out.xlsx"
        wb = Workbook()
        wb.active["A1"] = 1
        wb.save(source)
        original = source.read_bytes()
        args = {"file_path": str(source), "output_path": str(output), "changes": [{"target": "A1", "value": 42}], "mode": "dry_run"}
        response, preview = post(conn, rpc("tools/call", {"name": "office_patch", "arguments": args}), session)
        assert preview["result"]["structuredContent"]["changes_applied"] == 0
        assert not output.exists() and source.read_bytes() == original
        args["mode"] = "strict"
        response, result = post(conn, rpc("tools/call", {"name": "office_patch", "arguments": args}), session)
        assert not result["result"]["isError"]
        assert result["result"]["structuredContent"]["changes_applied"] == 1
        wb = load_workbook(output)
        assert wb.active["A1"].value == 42
        wb.close()
        assert source.read_bytes() == original
        denied, _ = post(conn, rpc("tools/list"), session, token="wrong")
        assert denied.status == 401
        bad_version, _ = post(conn, rpc("tools/list"), session, version="2024-11-05")
        assert bad_version.status == 400
        denied, _ = post(conn, rpc("tools/call", {"name": "restart_server"}), session)
        assert denied.status == 403
        response, _ = post(conn, {"jsonrpc": "2.0", "method": "notifications/initialized"}, session)
        assert response.status == 202
        conn.request("DELETE", "/mcp", headers={"Authorization": "Bearer " + TOKEN, "MCP-Protocol-Version": VERSION, "Mcp-Session-Id": session})
        response = conn.getresponse()
        response.read()
        assert response.status == 200
        response, _ = post(conn, rpc("tools/list"), session)
        assert response.status == 404
    finally:
        conn.close()


def test_session_sse_progress_and_single_stream(server):
    conn = http.client.HTTPConnection("127.0.0.1", server, timeout=5)
    stream_conn = http.client.HTTPConnection("127.0.0.1", server, timeout=5)
    conflict = http.client.HTTPConnection("127.0.0.1", server, timeout=5)
    try:
        session = initialize(conn)
        headers = {"Authorization": "Bearer " + TOKEN, "MCP-Protocol-Version": VERSION, "Mcp-Session-Id": session, "Accept": "text/event-stream"}
        stream_conn.request("GET", "/mcp", headers=headers)
        stream = stream_conn.getresponse()
        assert stream.status == 200
        assert stream.readline().startswith(b": connected")
        stream.readline()
        conflict.request("GET", "/mcp", headers=headers)
        response = conflict.getresponse()
        response.read()
        assert response.status == 409
        response, body = post(conn, rpc("tools/call", {"name": "office_help", "_meta": {"progressToken": "phase"}}), session)
        assert response.status == 200 and body["result"]["structuredContent"]["success"]
        events = []
        while len(events) < 2:
            line = stream.readline()
            assert line
            if line.startswith(b"data: "):
                events.append(json.loads(line[6:]))
        assert [e["params"]["progress"] for e in events] == [0, 1]
        assert all(e["method"] == "notifications/progress" for e in events)
        conn.request("DELETE", "/mcp", headers=headers)
        response = conn.getresponse()
        response.read()
        assert response.status == 200
        assert stream.read().strip() == b""
    finally:
        conn.close()
        stream_conn.close()
        conflict.close()


def test_http_negotiation_origin_size_and_header_guards(server):
    conn = http.client.HTTPConnection("127.0.0.1", server, timeout=5)
    try:
        for kwargs, expected in [({"version": None}, 400), ({"version": "future"}, 400), ({"Origin": "https://untrusted.example"}, 403), ({"Accept": "text/plain"}, 406)]:
            response, _ = post(conn, rpc("tools/list"), **kwargs)
            assert response.status == expected
        response, _ = post(conn, rpc("tools/list", {"padding": "x" * 70000}))
        assert response.status == 413
    finally:
        conn.close()
    with socket.create_connection(("127.0.0.1", server), timeout=5) as sock:
        sock.sendall(b"POST /mcp HTTP/1.1\r\nHost: localhost\r\nHost: duplicate\r\nContent-Length: 0\r\n\r\n")
        assert b"400" in sock.recv(4096).split(b"\r\n", 1)[0]


def test_cors_requires_explicit_allowed_origin(tmp_path):
    with running(tmp_path, "--http", "--allowed-origin", "https://review.example") as port:
        conn = http.client.HTTPConnection("127.0.0.1", port, timeout=5)
        try:
            conn.request("OPTIONS", "/mcp", headers={"Origin": "https://review.example", "Access-Control-Request-Method": "POST"})
            response = conn.getresponse()
            response.read()
            assert response.status == 204
            assert response.getheader("Access-Control-Allow-Origin") == "https://review.example"
            assert "Mcp-Session-Id" in response.getheader("Access-Control-Expose-Headers", "")
        finally:
            conn.close()


def test_legacy_sse_keeps_plain_port_selection_and_auth(tmp_path):
    with running(tmp_path) as port:
        conn = http.client.HTTPConnection("127.0.0.1", port, timeout=5)
        try:
            conn.request("GET", "/sse", headers={"Accept": "text/event-stream"})
            response = conn.getresponse()
            response.read()
            assert response.status == 401
        finally:
            conn.close()


def test_bind_policy_and_constant_time_token_hook(monkeypatch):
    monkeypatch.delenv("OFFICE_MCP_HTTP_TOKEN", raising=False)
    server = OfficeServer()
    with pytest.raises(ValueError, match="Non-loopback"):
        asyncio.run(server.run_streamable_http_async(host="0.0.0.0", port=0))
    monkeypatch.setenv("OFFICE_MCP_HTTP_TOKEN", TOKEN)
    server = OfficeServer()
    assert server.authenticate_request(method="POST", path="/mcp", headers={"authorization": "Bearer " + TOKEN}, peer=None)
    assert server.authenticate_request(method="POST", path="/mcp", headers={}, peer=None) is None
    with pytest.raises(ValueError, match="raw TCP"):
        asyncio.run(server.run_socket_async(host="0.0.0.0", port=0))


def test_raw_tcp_remains_loopback_legacy_compatible(tmp_path):
    with running(tmp_path, '--tcp') as port, socket.create_connection(('127.0.0.1', port), timeout=5) as sock:
        sock.sendall((json.dumps(rpc('initialize', {'protocolVersion': VERSION})) + '\n').encode())
        with sock.makefile('rb') as stream:
            result = json.loads(stream.readline())
            assert result['result']['serverInfo']['name'] == 'office-mcp-server'


def test_remote_exceptions_hide_internal_details(monkeypatch):
    from umcp_shared import MCPRequestContext

    async def check():
        server = OfficeServer()
        def broken():
            raise RuntimeError('sensitive-internal-path')
        server.register_tool('broken', broken)
        result = await server.process_request_async(json.dumps(rpc('tools/call', {'name': 'broken'})), context=MCPRequestContext(transport='streamable-http'))
        assert result['error']['code'] == -32603
        assert 'sensitive-internal-path' not in json.dumps(result)
    asyncio.run(check())


def test_session_expiry_releases_stream_and_subscriptions():
    from aioumcp import _AsyncStreamableHTTPSession

    async def check():
        server = OfficeServer()
        server.streamable_http_session_ttl_seconds = 1
        session = _AsyncStreamableHTTPSession('principal', VERSION, time.monotonic() - 5,
                                              time.monotonic() - 5, asyncio.Event(), asyncio.Queue())
        server._streamable_http_sessions['old'] = session
        server._resource_session_subscriptions['old'] = {'office://guidance/workflows'}
        server._expire_streamable_http_sessions(time.monotonic())
        assert not server._streamable_http_sessions
        assert not server._resource_session_subscriptions
        assert session.disconnect_event.is_set()
        assert session.queue.get_nowait() == b''
    asyncio.run(check())
