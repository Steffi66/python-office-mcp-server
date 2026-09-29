"""One canonical file-publication rollback; success and corruption controls earn no credit.

stage_patch discards a private staged copy on failure. It does not undo edits to
an in-memory OPC editor instance, and makes no such claim here.
"""

import zipfile
from pathlib import Path

from docx import Document

from tests.opc_preservation.test_preservation import (
    FIXTURE_ID, MAIN, OPAQUE, archive_parts, assert_graph, authored_source, edit_alpha, main_text,
)
from tests.opc_rollback.cases import CASE_KEY
from tools.mutation import stage_patch

ERROR = "injected failure after both staged parts changed"


def existing_destination(source, tmp_path):
    destination = tmp_path / "existing.docx"
    destination.write_bytes(source.read_bytes())
    document = Document(destination)
    document.paragraphs[0].text = "Prior"
    document.save(destination)
    names, parts = archive_parts(destination)
    assert main_text(parts) == ["Prior"]
    assert_graph(parts)
    return destination, names, parts, destination.read_bytes()


def change_two_parts(staged, names, original):
    """Edit real staged DOCX text, then change the opaque thumbnail's first byte."""
    result = edit_alpha(staged)
    path = Path(staged)
    entries = []
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            payload = archive.read(info)
            if info.filename == OPAQUE:
                payload = bytes((payload[0] ^ 1,)) + payload[1:]
            entries.append((info, payload))
    with zipfile.ZipFile(path, "w") as archive:
        for info, payload in entries:
            archive.writestr(info, payload)
    staged_names, staged_parts = archive_parts(path)
    assert staged_names == names and set(staged_parts) == set(original)
    assert main_text(staged_parts) == ["Beta"] and staged_parts[MAIN] != original[MAIN]
    assert staged_parts[OPAQUE] == bytes((original[OPAQUE][0] ^ 1,)) + original[OPAQUE][1:]
    assert len(staged_parts[OPAQUE]) == len(original[OPAQUE]) == 8324
    assert all(staged_parts[name] == original[name] for name in names if name not in {MAIN, OPAQUE})
    assert_graph(staged_parts)
    return result


def test_canonical_file_publication_rollback(rollback_case, tmp_path, request):
    case = rollback_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._opc_rollback_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    source = destination = None
    source_names = source_parts = destination_names = destination_parts = None
    source_bytes = destination_bytes = None
    staged_proof = []
    result = None
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                source, source_parts, sealed_parts = authored_source(tmp_path)
                source_bytes = source.read_bytes()
                source_names, observed_source = archive_parts(source)
                assert observed_source == source_parts and main_text(source_parts) == ["Alpha"]
                assert source_parts[OPAQUE] == sealed_parts[OPAQUE]
                destination, destination_names, destination_parts, destination_bytes = existing_destination(source, tmp_path)
                assert source_bytes != destination_bytes and destination_parts[MAIN] != source_parts[MAIN]
                assert source_names == destination_names and set(source_parts) == set(destination_parts)
                case["sourceProfile"] = {"sealedFixtureId": FIXTURE_ID, "authoredText": "Alpha",
                                         "priorDestinationText": "Prior", "members": len(source_names),
                                         "opaqueBytes": len(source_parts[OPAQUE])}
            elif index == 1:
                def fail_after_two_parts(staged):
                    change_two_parts(staged, source_names, source_parts)
                    staged_proof.append({"xmlText": "Beta", "opaqueChanged": True,
                                         "untouchedParts": len(source_names) - 2})
                    raise RuntimeError(ERROR)

                result = stage_patch(str(source), str(destination), "strict", 1, fail_after_two_parts)
                assert staged_proof == [{"xmlText": "Beta", "opaqueChanged": True, "untouchedParts": 15}]
                case["stagedBeforeFailure"] = staged_proof[0]
            else:
                assert result["success"] is False and result["status"] == "failed", result
                assert result["changes_applied"] == result["changes_planned"] == 0, result
                assert result["error"] == "Patch not committed: " + ERROR, result
                assert source.read_bytes() == source_bytes and destination.read_bytes() == destination_bytes
                actual_source_names, actual_source = archive_parts(source)
                actual_destination_names, actual_destination = archive_parts(destination)
                assert actual_source_names == source_names and actual_source == source_parts
                assert actual_destination_names == destination_names and actual_destination == destination_parts
                assert_graph(actual_source)
                assert_graph(actual_destination)
                assert [p.text for p in Document(source).paragraphs] == ["Alpha"]
                assert [p.text for p in Document(destination).paragraphs] == ["Prior"]
                assert not list(tmp_path.glob(".office-patch-*"))
                case["observedRollback"] = {"success": False, "changesApplied": 0,
                                             "sourceBytes": len(source_bytes),
                                             "existingDestinationBytes": len(destination_bytes),
                                             "sourcePartsIdentical": True, "destinationPartsIdentical": True,
                                             "stagingCleaned": True, "boundary": "private-file-publication"}
                # A second, successful invocation publishes BOTH parts. The failed
                # invocation's unchanged output cannot be explained by a no-op writer.
                published = tmp_path / "successful.docx"
                published.write_bytes(destination_bytes)
                successful_proof = []

                def publish_two_parts(staged):
                    edit_result = change_two_parts(staged, source_names, source_parts)
                    successful_proof.append(True)
                    return edit_result

                success = stage_patch(str(source), str(published), "strict", 1, publish_two_parts)
                assert successful_proof == [True]
                assert success["success"] is True and success["changes_planned"] == success["changes_applied"] == 1, success
                assert success["package_diff"]["added"] == success["package_diff"]["removed"] == []
                assert set(success["package_diff"]["changed"]) == {MAIN, OPAQUE}
                assert published.read_bytes() != destination_bytes
                published_names, saved = archive_parts(published)
                assert published_names == source_names and set(saved) == set(destination_parts)
                assert main_text(saved) == ["Beta"] and [p.text for p in Document(published).paragraphs] == ["Beta"]
                assert saved[OPAQUE] == bytes((source_parts[OPAQUE][0] ^ 1,)) + source_parts[OPAQUE][1:]
                assert all(saved[name] == source_parts[name] for name in source_names if name not in {MAIN, OPAQUE})
                assert_graph(saved)
                assert source.read_bytes() == source_bytes and destination.read_bytes() == destination_bytes
                assert archive_parts(source) == (source_names, source_parts)
                assert archive_parts(destination) == (destination_names, destination_parts)
                assert not list(tmp_path.glob(".office-patch-*"))
                case["positiveControl"] = {"publishedParts": sorted(success["package_diff"]["changed"]),
                                           "sourceAndPriorDestinationIdentical": True}
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()
