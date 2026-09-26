"""Enrolled legacy writers publish once, preserving signatures and read paths."""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from office_server import OfficeServer
from tools import mutation


def make(path):
    if path.suffix == ".xlsx":
        document = Workbook()
        document.active["A1"] = "target"
    elif path.suffix == ".docx":
        document = Document()
        document.add_paragraph("target")
    else:
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = "target"
    document.save(path)


@pytest.mark.parametrize("suffix,target", [(".docx", "target"), (".xlsx", "A1"), (".pptx", "1")])
def test_comment_preview_and_commit(tmp_path, suffix, target):
    source, output = tmp_path / ("source" + suffix), tmp_path / ("output" + suffix)
    make(source)
    before = source.read_bytes()
    output.write_bytes(b"retain old destination")
    server = OfficeServer()
    args = {"file_path": str(source), "operation": "add", "target": target, "text": "review", "output_path": str(output)}
    preview = server.tool_office_comment(**args, mode="dry_run")
    assert preview["success"], preview
    assert preview["changes_applied"] == 0
    assert source.read_bytes() == before and output.read_bytes() == b"retain old destination"
    result = server.tool_office_comment(**args, mode="safe")
    assert result["changes_applied"] == 1, result
    assert source.read_bytes() == before
    assert "review" in str(server.tool_office_comment(str(output), operation="get"))


@pytest.mark.parametrize("mode", ["dry_run", "strict", "safe"])
def test_word_table_uses_top_level_output(tmp_path, mode):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Name"
    table.cell(1, 0).text = "old"
    doc.save(source)
    before = source.read_bytes()
    result = OfficeServer().tool_office_table(str(source), operation="add_row", table_id="0",
                                             data={"Name": "new"}, output_path=str(output), mode=mode)
    assert result["success"], result
    assert source.read_bytes() == before
    if mode == "dry_run":
        assert not output.exists()
    else:
        assert len(Document(output).tables[0].rows) == 3


def test_specialised_writer_failure_does_not_publish(tmp_path, monkeypatch):
    source, output = tmp_path / "source.pptx", tmp_path / "out.pptx"
    make(source)
    before = source.read_bytes()
    output.write_bytes(b"previous")

    def fail(_path):
        raise ValueError("injected validation")

    monkeypatch.setattr(mutation, "validate_staged_document", fail)
    result = OfficeServer().tool_pptx_add_slide(str(source), output_path=str(output))
    assert not result["success"] and "injected validation" in result["error"]
    assert source.read_bytes() == before
    assert output.read_bytes() == b"previous"


def test_nested_generic_writer_publishes_once(tmp_path, monkeypatch):
    source = tmp_path / "source.xlsx"
    make(source)
    server = OfficeServer()
    server.tool_office_comment(str(source), operation="add", target="A1", text="review")
    original = mutation.os.replace
    publications = []

    def replace(src, dst):
        if Path(dst) == source:
            publications.append(str(dst))
        return original(src, dst)

    monkeypatch.setattr(mutation.os, "replace", replace)
    result = server.tool_office_comment(str(source), operation="delete", target="A1")
    assert result["success"], result
    assert publications == [str(source)]
    wb = load_workbook(source)
    try:
        assert wb.active["A1"].comment is None
    finally:
        wb.close()
