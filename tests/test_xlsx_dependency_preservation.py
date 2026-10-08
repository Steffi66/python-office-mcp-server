"""Cell edits must preserve package dependencies and invalidate stale calculation data."""

import shutil
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from office_server import OfficeServer
from tests.fixture_paths import shared_fixture
from tools.xlsx_preservation import merge_styles

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def parts(path):
    with zipfile.ZipFile(path) as archive:
        return {n: archive.read(n) for n in archive.namelist()}


def test_multiline_edit_saves_style_dependency_and_retains_opaque_parts(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "out.xlsx"
    shutil.copy2(shared_fixture("default-style.xlsx"), source)
    before = parts(source)
    result = OfficeServer().tool_office_patch(
        str(source), [{"target": "A1", "value": "first\nsecond"}],
        mode="safe", output_path=str(output),
    )
    assert result["changes_applied"] == 1, result
    after = parts(output)
    assert after.keys() == before.keys()
    for name in before:
        if name not in {"xl/worksheets/sheet1.xml", "xl/styles.xml"}:
            assert after[name] == before[name], name
    workbook = load_workbook(output)
    try:
        assert workbook.active["A1"].value == "first\nsecond"
        assert workbook.active["A1"].alignment.wrap_text
    finally:
        workbook.close()
    styles = ET.fromstring(after["xl/styles.xml"])
    count = len(styles.find(S + "cellXfs"))
    for c in ET.fromstring(after["xl/worksheets/sheet1.xml"]).iter(S + "c"):
        assert 0 <= int(c.get("s", "0")) < count


def test_cross_sheet_formula_cache_is_invalidated_without_calculation(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "out.xlsx"
    shutil.copy2(shared_fixture("cross-sheet-cache.xlsx"), source)
    before = parts(source)
    result = OfficeServer().tool_office_patch(
        str(source), [{"target": "Input!A1", "value": 10}],
        mode="safe", output_path=str(output),
    )
    assert result["changes_applied"] == 1, result
    assert result["calculation_state"] == "recalculation-required"
    assert result["preservation"]["cache_policy"] == "invalidate-all-formula-caches"
    for data_only, expected in [(True, None), (False, "=Input!A1*2")]:
        wb = load_workbook(output, data_only=data_only)
        try:
            assert wb["Input"]["A1"].value == 10
            assert wb["Calc"]["A1"].value == expected
        finally:
            wb.close()
    after = parts(output)
    for name in before:
        if name not in {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/workbook.xml"}:
            assert before[name] == after[name], name
    props = ET.fromstring(after["xl/workbook.xml"]).find(S + "calcPr")
    assert props.get("forceFullCalc") == "1"


def test_existing_custom_style_indices_remain_valid(tmp_path):
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active["B1"] = 1.25
    wb.active["B1"].number_format = '#,##0.0000" units"'
    wb.active["B1"].alignment = Alignment(horizontal="right")
    wb.save(source)
    result = OfficeServer().tool_office_patch(str(source), [{"target": "A1", "value": "a\nb"}])
    assert result["changes_applied"] == 1, result
    wb = load_workbook(source)
    try:
        assert wb.active["B1"].number_format == '#,##0.0000" units"'
        assert wb.active["B1"].alignment.horizontal == "right"
        assert wb.active["A1"].alignment.wrap_text
    finally:
        wb.close()


def test_original_styles_xml_is_preserved_for_cell_edit(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "out.xlsx"
    shutil.copy2(shared_fixture("default-style.xlsx"), source)

    before = parts(source)

    result = OfficeServer().tool_office_patch(
        str(source),
        [{"target": "A1", "value": "changed"}],
        mode="safe",
        output_path=str(output),
    )

    assert result["changes_applied"] == 1, result

    after = parts(output)

    assert after["xl/styles.xml"] == before["xl/styles.xml"]

def test_style_reindexing_is_refused():
    original = f'<styleSheet xmlns="{S[1:-1]}"><cellXfs count="1"><xf fontId="0"/></cellXfs></styleSheet>'.encode()
    rewritten = original.replace(b'fontId="0"', b'fontId="1"')
    with pytest.raises(ValueError, match="registry rewrite"):
        merge_styles(original, rewritten)


def test_calculation_chain_relationship_and_content_type_are_removed(tmp_path):
    source = tmp_path / "source.xlsx"
    entries = parts(shared_fixture("cross-sheet-cache.xlsx"))
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    rels = ET.fromstring(entries["xl/_rels/workbook.xml.rels"])
    ET.SubElement(rels, rel_ns + "Relationship", {
        "Id": "rIdCalcChain", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
        "Target": "calcChain.xml",
    })
    entries["xl/_rels/workbook.xml.rels"] = ET.tostring(rels)
    types = ET.fromstring(entries["[Content_Types].xml"])
    ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
        "PartName": "/xl/calcChain.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
    })
    entries["[Content_Types].xml"] = ET.tostring(types)
    entries["xl/calcChain.xml"] = f'<calcChain xmlns="{S[1:-1]}"><c r="A1" i="2"/></calcChain>'.encode()
    with zipfile.ZipFile(source, "w") as z:
        for name, payload in entries.items():
            z.writestr(name, payload)
    result = OfficeServer().tool_office_patch(str(source), [{"target": "Input!A1", "value": 10}])
    assert result["changes_applied"] == 1, result
    after = parts(source)
    assert "xl/calcChain.xml" not in after
    assert b"calcChain" not in after["xl/_rels/workbook.xml.rels"]
    assert b"calcChain" not in after["[Content_Types].xml"]
