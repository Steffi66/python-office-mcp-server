"""Execute the eight authored comment metadata and filter predicates without vacuity."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_filter_predicates.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _rows(reader, path, selected, count, *, author=None):
    result = reader.tool_word_get_comments(str(path), filter=selected, author=author)
    assert "error" not in result and result.get("filter") == selected
    assert result["comment_count"] == len(result["comments"]) == count
    return result["comments"]


def _author_two(path, authoring):
    document = Document()
    document.add_paragraph("Alpha target")
    document.add_paragraph("Beta target")
    document.save(path)
    assert [p.text for p in Document(path).paragraphs if p.text] == ["Alpha target", "Beta target"]
    for target, text, author in (("Alpha target", "Comment alpha", "Manuel"),
                                 ("Beta target", "Comment beta", "Rui Carmo")):
        result = authoring.tool_word_add_comment(
            file_path=str(path), target_text=target, comment_text=text, author=author)
        assert result.get("success") is True


def _initial(path, reader):
    rows = _rows(reader, path, "all", 2)
    assert [(row["text"], row["author"], row["done"]) for row in rows] == [
        ("Comment alpha", "Manuel", False), ("Comment beta", "Rui Carmo", False)]
    first, second = (str(row["id"]) for row in rows)
    assert first != second
    body = etree.fromstring(_parts(path)["word/document.xml"])
    targets = [node for node in body.findall(".//" + W + "body/" + W + "p")
               if "".join(node.itertext()) in ("Alpha target", "Beta target")]
    assert ["".join(node.itertext()) for node in targets] == ["Alpha target", "Beta target"]
    for target, expected in zip(targets, (first, second)):
        assert [node.get(W + "id") for node in target.findall(".//" + W + "commentRangeStart")] == [expected]
        assert [node.get(W + "id") for node in target.findall(".//" + W + "commentRangeEnd")] == [expected]
    return rows, first, second


def test_canonical_filter_predicates(filter_predicates_case, tmp_path, request):
    case = filter_predicates_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_filter_predicates_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "filter-predicates.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                _author_two(path, authoring)
            elif index == 1:
                initial, first, second = _initial(path, reader)
                source = _parts(path)
                response = reader.tool_word_resolve_comment(str(path), first, True)
                assert path.is_file()
                case["firstId"] = first
                case["secondId"] = second
            elif index == 2:
                assert {"done", "is_reply", "parent_id", "para_id"}.issubset(initial[0])
                assert initial[0]["done"] is False and initial[0]["is_reply"] is False
                assert initial[0]["parent_id"] is None
            elif index == 3:
                assert response.get("success") is True and response.get("done") is True
                assert str(response.get("comment_id")) == first
                assert str(response.get("thread_root_comment_id")) == first
            elif index == 4:
                open_rows = _rows(reader, path, "open", 1)
                resolved_rows = _rows(reader, path, "resolved", 1)
                mine_rows = _rows(reader, path, "mine", 1, author="rUi CaRmO")
                assert [(str(row["id"]), row["author"], row["done"]) for row in open_rows] == [
                    (second, "Rui Carmo", False)]
                assert [(str(row["id"]), row["author"], row["done"]) for row in resolved_rows] == [
                    (first, "Manuel", True)]
                assert [(str(row["id"]), row["author"], row["done"]) for row in mine_rows] == [
                    (second, "Rui Carmo", False)]
                all_after = _rows(reader, path, "all", 2)
                assert [(str(row["id"]), row["done"]) for row in all_after] == [(first, True), (second, False)]
                assert _parts(path)["word/document.xml"] == source["word/document.xml"]
            elif index == 5:
                assert open_rows and all(row["done"] is False for row in open_rows)
                assert {str(row["id"]) for row in open_rows} == {second}
            elif index == 6:
                assert resolved_rows and all(row["done"] is True for row in resolved_rows)
                assert {str(row["id"]) for row in resolved_rows} == {first}
            else:
                assert mine_rows and all(row["author"].casefold() == "rui carmo" for row in mine_rows)
                assert {str(row["id"]) for row in mine_rows} == {second}
                case["filtersNonempty"] = True
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_opposite_resolve_flips_nonempty_filter_membership(tmp_path):
    path = tmp_path / "opposite.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    _author_two(path, authoring)
    _, first, second = _initial(path, reader)
    assert _rows(reader, path, "resolved", 0) == []
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "open", 2)] == [
        (first, False), (second, False)]
    response = reader.tool_word_resolve_comment(str(path), second, True)
    assert response.get("success") is True and str(response.get("comment_id")) == second
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "resolved", 1)] == [(second, True)]
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "open", 1)] == [(first, False)]
    mine = _rows(reader, path, "mine", 1, author="  RUI CARMO  ")
    assert [(str(row["id"]), row["done"]) for row in mine] == [(second, True)]
    assert _rows(reader, path, "mine", 0, author="Nobody") == []


def test_missing_id_does_not_produce_resolved_rows(tmp_path):
    path = tmp_path / "missing.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    _author_two(path, authoring)
    _, first, second = _initial(path, reader)
    source = path.read_bytes()
    response = reader.tool_word_resolve_comment(str(path), str(max(int(first), int(second)) + 100), True)
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source
    assert _rows(reader, path, "resolved", 0) == []
    assert [(str(row["id"]), row["done"]) for row in _rows(reader, path, "open", 2)] == [
        (first, False), (second, False)]
