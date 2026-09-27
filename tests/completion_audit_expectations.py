"""Exact native expectations for the planned read-only Word completion audit.

Author inputs from the sealed blank DOCX; assert literal responses independently of
what the audit implementation returns. This is not a shared Gherkin binding.
"""

import hashlib
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from tests.fixture_paths import fixture_path

BLANK_ID = "fixture-b051c0c2ff43f2ab9213e19a52ccbc51217537cf3b336d54a68a91beb6670f9a"
EMPTY_ISSUES = {
    "placeholders": [], "empty_sections": ["Document Start"],
    "empty_table_cells": [], "instruction_remnants": [], "pending_track_changes": False,
}


def snapshot(path: Path):
    data = path.read_bytes()
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    return hashlib.sha256(data).hexdigest(), data, members


def assert_unchanged(path: Path, before):
    assert snapshot(path) == before


def blank_source():
    path = fixture_path(BLANK_ID)
    before = snapshot(path)
    assert before[0] == BLANK_ID.removeprefix("fixture-")
    document = Document(path)
    assert len(document.tables) == 0
    assert [(p.text, p.style.name) for p in document.paragraphs] == [("", "Normal")]
    return path, before


def authored_source(tmp_path: Path, name: str, heading: str, body: str, blank: Path):
    document = Document(blank)
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    path = tmp_path / name
    document.save(path)
    reopened = Document(path)
    assert [(p.text, p.style.name) for p in reopened.paragraphs][-2:] == [
        (heading, "Heading 1"), (body, "Normal"),
    ]
    return path, snapshot(path)


def assert_audit(tools, path: Path, before, blank: Path, blank_before, expected):
    assert tools.tool_word_audit_completion(str(path)) == {"file": str(path), **expected}
    assert_unchanged(path, before)
    assert_unchanged(blank, blank_before)


def assert_complete_audit(tools, tmp_path: Path, *, heading: str, body: str):
    blank, blank_before = blank_source()
    path, before = authored_source(tmp_path, "complete.docx", heading, body, blank)
    assert_audit(tools, path, before, blank, blank_before, {
        "success": True, "status": "READY", "score": 95,
        "recommendation": "Document appears complete and ready for review.",
        "summary": {"placeholders_found": 0, "empty_sections": 1,
                    "empty_table_cells": 0, "instruction_remnants": 0,
                    "pending_changes": False},
        "issues": EMPTY_ISSUES,
        "next_tools": ["word_cleanup_sow"],
    })


def assert_placeholder_and_missing_audit(tools, tmp_path: Path):
    blank, blank_before = blank_source()
    path, before = authored_source(tmp_path, "placeholder-audit.docx", "Project: <Name>", "Customer: [TBD]", blank)
    assert_audit(tools, path, before, blank, blank_before, {
        "success": True, "status": "NEEDS_REVIEW", "score": 85,
        "recommendation": "Minor issues found. Review flagged items before finalizing.",
        "summary": {"placeholders_found": 2, "empty_sections": 1,
                    "empty_table_cells": 0, "instruction_remnants": 0,
                    "pending_changes": False},
        "issues": {**EMPTY_ISSUES, "placeholders": [
            {"text": "<Name>", "location": "Project: <Name>", "context": "Project: <Name>"},
            {"text": "[TBD]", "location": "Project: <Name>", "context": "Customer: [TBD]"},
        ]},
        "next_tools": ["word_audit_sow", "word_fix_split_placeholders", "word_patch_placeholder"],
    })
    missing = tmp_path / "missing-completion-audit.docx"
    assert not missing.exists()
    assert tools.tool_word_audit_completion(str(missing)) == {"error": f"File not found: {missing}"}
    assert not missing.exists()
    assert_unchanged(path, before)
    assert_unchanged(blank, blank_before)
