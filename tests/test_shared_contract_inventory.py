"""Validate pinned acceptance inputs; planned scenarios are not execution passes."""

import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from tests.acceptance.ledger import inventory
from tests.fixture_paths import (
    CONTRACT,
    FEATURE,
    FIXTURE_SOURCE,
    MUTATION_SCENARIO_IDS,
    acceptance_mapping_path,
    mutation_feature_paths,
    REPOSITORY,
    fixture_path,
    load_fixture_assets,
    mutation_contract,
    preserved_members,
    shared_fixture,
    template_asset_ids,
    verified_metadata,
)

POLICY = mutation_contract()


def digest(data):
    return hashlib.sha256(data).hexdigest()


def test_shared_pack_matches_pinned_manifest():
    assert not (FIXTURE_SOURCE / "shared").exists()
    assert verified_metadata("contracts/mutation-safety.json", "workflow-contract") == CONTRACT.read_bytes()
    for path in mutation_feature_paths(POLICY):
        assert verified_metadata(path.relative_to(FIXTURE_SOURCE).as_posix(), "workflow") == path.read_bytes()
    assert POLICY["fixturePolicy"] == {"membership": "exact", "preserve": "all-except-allowed"}


def test_expanded_inventory_has_unique_ids_and_exact_fixture_pins():
    paths = mutation_feature_paths(POLICY)
    compiled = inventory(paths, scenario_ids=MUTATION_SCENARIO_IDS)  # Official Gherkin compiler.
    assert len(compiled) == POLICY["expandedCaseCount"] == 19
    assert {c["scenarioId"] for c in compiled} == set(POLICY["scenarioIds"]) == set(MUTATION_SCENARIO_IDS)
    mapping = json.loads(acceptance_mapping_path(POLICY).read_text())
    assert {c["stableCaseKey"] for c in compiled} == set(mapping["implementedCaseKeys"])
    fixtures = {f["id"]: f for f in POLICY["fixtures"]}
    for case in compiled:
        assert case["outcome"] == "planned"
        assert case["featureSha256"] == digest(Path(case["feature"]).read_bytes())
        assert any(step["type"] == "Outcome" for step in case["steps"])
        fixture_id = next(re.fullmatch(r'fixture "([^"]+)" verified against the fixture manifest', s["text"])[1]
                          for s in case["steps"] if s["text"].startswith('fixture "'))
        record = fixtures[fixture_id]
        assert "fixture-" + digest(shared_fixture(fixture_id).read_bytes()) == record["assetId"]


@pytest.mark.parametrize("fixture", POLICY["fixtures"], ids=lambda f: f["id"])
def test_fixture_and_member_hashes(fixture):
    path = shared_fixture(fixture["id"])
    asset = load_fixture_assets(FIXTURE_SOURCE)[fixture["assetId"]]
    assert path == FIXTURE_SOURCE / asset["path"]
    assert digest(path.read_bytes()) == asset["sha256"]
    assert any(origin.get("revision") == "36ac406ad9d4bd3e7538b4bcc7aa2fb0e51cc943" for origin in asset["origins"])
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names))
        assert archive.testzip() is None
        assert set(names) == set(fixture["memberSha256"])
        for name in names:
            assert digest(archive.read(name)) == fixture["memberSha256"][name]
        assert "customXml/preservation-sentinel.xml" in preserved_members(fixture)
        assert set(preserved_members(fixture)) == set(names) - set(fixture["allowedChangedPartsForSuccess"])


def test_pinned_word_and_slide_facts():
    document = Document(shared_fixture("present-placeholder.docx"))
    assert [p.text for p in document.paragraphs] == ["<Present>"]
    presentation = Presentation(shared_fixture("title-and-subtitle.pptx"))
    assert len(presentation.slides) == 1
    assert presentation.slides[0].shapes.title.text == "Original title"
    assert presentation.slides[0].placeholders[1].text == "Original subtitle"


def test_pinned_workbook_style_and_cache_facts():
    path = shared_fixture("default-style.xlsx")
    workbook = load_workbook(path)
    try:
        assert workbook.active["A1"].value == "before"
        assert "Missing" not in workbook.sheetnames
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        styles = ET.fromstring(archive.read("xl/styles.xml"))
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        assert len(styles.find(ns + "cellXfs")) == 1

    for data_only, expected in ((True, 2), (False, "=Input!A1*2")):
        workbook = load_workbook(shared_fixture("cross-sheet-cache.xlsx"), data_only=data_only)
        try:
            assert workbook["Input"]["A1"].value == 1
            assert workbook["Calc"]["A1"].value == expected
        finally:
            workbook.close()


def test_central_fixture_manifest_matches_python_inputs():
    records = load_fixture_assets(FIXTURE_SOURCE)
    mapping = template_asset_ids()
    assert len(mapping) == 37  # Two defaults and 35 document fixtures; notices stay metadata.
    prefix = "fixtures/python-office-mcp-server/tests/_templates/"
    expected = {alias[len(prefix):]: record["id"] for record in records.values()
                for alias in record["aliases"] if alias.startswith(prefix)}
    assert mapping == expected
    for asset_id in mapping.values():
        assert fixture_path(asset_id).is_file()
    actual_files = {p.relative_to(FIXTURE_SOURCE).as_posix()
                    for p in (FIXTURE_SOURCE / "fixtures").rglob("*") if p.is_file()}
    assert actual_files == {record["path"] for record in records.values()}
    for asset_id in records:
        fixture_path(asset_id)  # Verify one physical payload per unique ID/hash.


def test_python_constants_match_selected_shared_facts():
    from tools import word_tools

    bindings = {
        "namespaces": {"NSWordprocessingML": word_tools.W_NS, "NSWord14": word_tools.W14_NS,
                       "NSContentTypes": word_tools.PKG_CT_NS, "NSRelationships": word_tools.PKG_REL_NS},
        "relationships": {"RelTypeCommentsExtended": word_tools.REL_COMMENTS_EXTENDED},
        "content-types": {"ContentTypeCommentsExtendedSpecified": word_tools.CT_COMMENTS_EXTENDED},
    }
    for group, constants in bindings.items():
        catalog = json.loads((FIXTURE_SOURCE / "facts" / (group + ".json")).read_text())
        assert catalog["schemaVersion"] == 1
        values = {row["id"]: row for row in catalog["values"]}
        assert len(values) == len(catalog["values"])
        for name, value in constants.items():
            assert values[name]["status"] in {"specified", "observed"}
            assert values[name]["value"] == value
    # A documented alias does not silently become an accepted runtime constant.
    content_types = json.loads((FIXTURE_SOURCE / "facts/content-types.json").read_text())["values"]
    disputed = next(row for row in content_types if row["id"] == "ContentTypeCommentsExtended")
    assert disputed["status"] == "disputed"
    assert disputed["value"] != word_tools.CT_COMMENTS_EXTENDED


def test_fixture_submodule_matches_common_tag_and_seals():
    from tests.fixture_paths import verify_selected_fixture_source

    verify_selected_fixture_source()
