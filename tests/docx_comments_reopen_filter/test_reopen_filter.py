"""Execute one authored comment reopen-filter case; controls earn no sibling credit."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_reopen_filter.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _comments(response, selected_filter, count):
    assert "error" not in response and response.get("filter") == selected_filter
    rows = response["comments"]
    assert response["comment_count"] == len(rows) == count
    return rows


def _author_two(path, authoring):
    doc = Document()
    doc.add_paragraph("Alpha target")
    doc.add_paragraph("Beta target")
    doc.save(path)
    assert path.is_file() and [p.text for p in Document(path).paragraphs if p.text] == ["Alpha target", "Beta target"]
    for target, body, author in [("Alpha target", "Comment alpha", "Manuel"),
                                 ("Beta target", "Comment beta", "Rui Carmo")]:
        response = authoring.tool_word_add_comment(
            file_path=str(path), target_text=target, comment_text=body, author=author)
        assert response.get("success") is True


def _ids_and_anchors(path, reader):
    rows = _comments(reader.tool_word_get_comments(str(path), filter="all"), "all", 2)
    assert [(r["text"], r["author"], r["done"]) for r in rows] == [
        ("Comment alpha", "Manuel", False), ("Comment beta", "Rui Carmo", False)]
    first_id, second_id = (str(r["id"]) for r in rows)
    assert first_id != second_id
    body = etree.fromstring(_parts(path)["word/document.xml"])
    paragraphs = body.findall(".//" + W + "body/" + W + "p")
    targets = [p for p in paragraphs if "".join(p.itertext()) in ("Alpha target", "Beta target")]
    assert ["".join(p.itertext()) for p in targets] == ["Alpha target", "Beta target"]
    for target, expected_id in zip(targets, (first_id, second_id)):
        assert [n.get(W + "id") for n in target.findall(".//" + W + "commentRangeStart")] == [expected_id]
        assert [n.get(W + "id") for n in target.findall(".//" + W + "commentRangeEnd")] == [expected_id]
    return first_id, second_id


def test_canonical_authored_reopen_filter(reopen_filter_case, tmp_path, request):
    case = reopen_filter_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_reopen_filter_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "reopen-filter.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                _author_two(path, authoring)
            elif index == 1:
                first_id, second_id = _ids_and_anchors(path, reader)
                original_parts = _parts(path)
                case["selectedId"] = first_id
            elif index == 2:
                resolved_response = reader.tool_word_resolve_comment(str(path), first_id, True)
                assert path.is_file()
            elif index == 3:
                assert resolved_response.get("success") is True
                assert resolved_response.get("done") is True
                assert str(resolved_response.get("thread_root_comment_id")) == first_id
            elif index == 4:
                filtered = _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 1)
                assert [(str(r["id"]), r["done"]) for r in filtered] == [(first_id, True)]
                all_rows = _comments(reader.tool_word_get_comments(str(path), filter="all"), "all", 2)
                assert [(str(r["id"]), r["done"]) for r in all_rows] == [(first_id, True), (second_id, False)]
                case["resolvedReadCount"] = len(filtered)
            elif index == 5:
                reopened_response = reader.tool_word_resolve_comment(str(path), first_id, False)
                assert path.is_file()
            elif index == 6:
                assert reopened_response.get("success") is True and reopened_response.get("done") is False
                assert str(reopened_response.get("thread_root_comment_id")) == first_id
            else:
                reopened = _comments(reader.tool_word_get_comments(str(path), filter="open"), "open", 2)
                assert [(str(r["id"]), r["text"], r["author"], r["done"]) for r in reopened] == [
                    (first_id, "Comment alpha", "Manuel", False),
                    (second_id, "Comment beta", "Rui Carmo", False)]
                assert _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 0) == []
                restored_parts = _parts(path)
                old_comments = etree.fromstring(original_parts["word/comments.xml"])
                new_comments = etree.fromstring(restored_parts["word/comments.xml"])
                assert [(n.get(W + "id"), "".join(n.itertext())) for n in old_comments] == [
                    (n.get(W + "id"), "".join(n.itertext())) for n in new_comments]
                extension = etree.fromstring(restored_parts["word/commentsExtended.xml"])
                assert any(node.get(W15 + "paraId") == str(reopened[0]["para_id"])
                           and node.get(W15 + "done") == "0" for node in extension.findall(W15 + "commentEx"))
                case["observedReopenedId"] = first_id
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_opposite_id_noop_and_filter_controls(tmp_path):
    path = tmp_path / "opposite.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    _author_two(path, authoring)
    first_id, second_id = _ids_and_anchors(path, reader)
    assert _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 0) == []
    assert reader.tool_word_resolve_comment(str(path), second_id, True).get("success") is True
    assert [(str(r["id"]), r["done"]) for r in _comments(
        reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 1)] == [(second_id, True)]
    assert [(str(r["id"]), r["done"]) for r in _comments(
        reader.tool_word_get_comments(str(path), filter="open"), "open", 1)] == [(first_id, False)]
    assert reader.tool_word_resolve_comment(str(path), second_id, False).get("success") is True
    prior = path.read_bytes()
    again = reader.tool_word_resolve_comment(str(path), second_id, False)
    assert again.get("success") is True and again.get("unchanged") is True
    assert path.read_bytes() == prior
    assert [(str(r["id"]), r["done"]) for r in _comments(
        reader.tool_word_get_comments(str(path), filter="open"), "open", 2)] == [(first_id, False), (second_id, False)]


def test_missing_id_does_not_publish(tmp_path):
    path = tmp_path / "missing.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    _author_two(path, authoring)
    first_id, second_id = _ids_and_anchors(path, reader)
    source = path.read_bytes()
    unknown = str(max(int(first_id), int(second_id)) + 100)
    response = reader.tool_word_resolve_comment(str(path), unknown, False)
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source
    assert _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 0) == []
