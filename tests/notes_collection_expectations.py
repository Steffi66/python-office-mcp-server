"""Literal outcomes for the planned four-slide PPTX notes-collection profile.

Input construction is test setup from a sealed blank presentation. Assertions
use independent expected values; neither the reader nor a saved output supplies
its own oracle.
"""

import hashlib
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation

from tests.fixture_paths import fixture_path

BLANK_FIXTURE_ID = "fixture-c54a7b746c0328fc1930525edd91387eedbbca69a6f388f7ec024150187b6bab"
NOTES = ("Notes for slide 1", "Notes for slide 2", "Notes for slide 3")


def snapshot(path: Path):
    data = path.read_bytes()
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    return hashlib.sha256(data).hexdigest(), data, members


def assert_unchanged(path: Path, before):
    assert snapshot(path) == before


def assert_no_fourth_notes(path: Path):
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        assert "ppt/notesSlides/notesSlide4.xml" not in names
        relationships = archive.read("ppt/slides/_rels/slide4.xml.rels") if "ppt/slides/_rels/slide4.xml.rels" in names else b""
        assert b"/notesSlide" not in relationships
    assert not Presentation(path).slides[3].has_notes_slide


def prepare_notes_source(tmp_path: Path, *, titled: bool):
    blank = fixture_path(BLANK_FIXTURE_ID)
    blank_before = snapshot(blank)
    assert blank_before[0] == BLANK_FIXTURE_ID.removeprefix("fixture-")
    presentation = Presentation(blank)
    assert len(presentation.slides) == 0
    for index in range(4):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        if titled:
            slide.shapes.title.text = f"Slide {index + 1}"
        if index < 3:
            slide.notes_slide.notes_text_frame.text = NOTES[index]
        else:
            assert not slide.has_notes_slide
    path = tmp_path / "notes-collection.pptx"
    presentation.save(path)
    reopened = Presentation(path)
    assert len(reopened.slides) == 4
    assert [slide.shapes.title.text for slide in reopened.slides] == (
        [f"Slide {index}" for index in range(1, 5)] if titled else [""] * 4
    )
    assert [slide.has_notes_slide for slide in reopened.slides] == [True, True, True, False]
    assert_no_fourth_notes(path)
    assert_unchanged(blank, blank_before)
    return path, snapshot(path), blank, blank_before


def assert_notes_collection(tools, tmp_path: Path, *, titled: bool):
    path, before, blank, blank_before = prepare_notes_source(tmp_path, titled=titled)
    expected_all = {
        "file": path.name,
        "slides_with_notes": 3,
        "total_slides": 4,
        "notes": [{"slide_number": n, "notes": text} for n, text in enumerate(NOTES, 1)],
    }

    def check(result, expected):
        assert result == expected
        assert_unchanged(path, before)
        assert_no_fourth_notes(path)
        assert_unchanged(blank, blank_before)

    check(tools.tool_pptx_get_notes(str(path)), expected_all)
    for number, text in enumerate(NOTES, 1):
        check(tools.tool_pptx_get_notes(str(path), slide_number=number), {
            "file": path.name, "slide_number": number, "has_notes": True, "notes": text,
        })
    check(tools.tool_pptx_get_notes(str(path), slide_number=4), {
        "file": path.name, "slide_number": 4, "has_notes": False, "notes": "",
    })
    for number in (0, 5):
        check(tools.tool_pptx_get_notes(str(path), slide_number=number), {
            "error": f"Slide {number} not found. Presentation has 4 slides.",
        })
