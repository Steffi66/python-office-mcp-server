"""One staged DOCX edit, one unrelated opaque payload; sibling controls earn no credit."""

import hashlib
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from tests.fixture_paths import fixture_path
from tests.opc_preservation.cases import CASE_KEY
from tools.mutation import stage_patch

FIXTURE_ID = "fixture-d9d6a313182a71a73d75a26a0ff3b7826dbd2e300e1d202114ec9f8fb018fda5"
MAIN = "word/document.xml"
OPAQUE = "docProps/thumbnail.jpeg"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
THUMBNAIL_SHA256 = "96367138dc44ce09bf2c8f0f8e49348a1478d2c5c0af69bbc2bbc38b63cdcead"


def archive_parts(path):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        names = [info.filename for info in infos]
        assert len(names) == len(set(names)) == 17
        assert MAIN in names and OPAQUE in names
        parts = {name: archive.read(name) for name in names}
    return names, parts


def assert_graph(parts):
    rels = etree.fromstring(parts["_rels/.rels"])
    assert rels.tag == f"{{{PACKAGE_REL}}}Relationships"
    root_parts = [r for r in rels if r.tag == f"{{{PACKAGE_REL}}}Relationship"
                  and r.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"]
    thumbnails = [r for r in rels if r.tag == f"{{{PACKAGE_REL}}}Relationship"
                  and r.get("Type") == "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail"]
    assert len(root_parts) == 1 and root_parts[0].get("Target") == MAIN
    assert len(thumbnails) == 1 and thumbnails[0].get("Target") == OPAQUE
    types = etree.fromstring(parts["[Content_Types].xml"])
    assert any(node.get("Extension") == "jpeg" and node.get("ContentType") == "image/jpeg"
               for node in types if node.tag.endswith("}Default"))
    assert any(node.get("PartName") == "/" + MAIN and
               node.get("ContentType") == "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
               for node in types if node.tag.endswith("}Override"))
    assert len(parts[OPAQUE]) == 8324


def main_text(parts):
    root = etree.fromstring(parts[MAIN])
    assert root.tag == f"{{{W}}}document"
    return [node.text for node in root.iter(f"{{{W}}}t")]


def authored_source(tmp_path):
    sealed = fixture_path(FIXTURE_ID)
    sealed_names, sealed_parts = archive_parts(sealed)
    assert_graph(sealed_parts)
    assert hashlib.sha256(sealed_parts[OPAQUE]).hexdigest() == THUMBNAIL_SHA256  # sealed-asset provenance
    assert main_text(sealed_parts) == []  # source is blank; do not pretend it contains Alpha
    source = tmp_path / "source.docx"
    root = etree.fromstring(sealed_parts[MAIN])
    body = root.find(f"{{{W}}}body")
    assert body is not None and [etree.QName(node).localname for node in body] == ["sectPr"]
    paragraph = etree.Element(f"{{{W}}}p")
    etree.SubElement(etree.SubElement(paragraph, f"{{{W}}}r"), f"{{{W}}}t").text = "Alpha"
    body.insert(0, paragraph)
    edited = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
    assert etree.fromstring(edited) is not None and edited != sealed_parts[MAIN]
    with zipfile.ZipFile(sealed) as original, zipfile.ZipFile(source, "w") as archive:
        for info in original.infolist():
            archive.writestr(info, edited if info.filename == MAIN else original.read(info))
    names, parts = archive_parts(source)
    assert names == sealed_names and main_text(parts) == ["Alpha"]
    assert_graph(parts)
    assert all(parts[name] == sealed_parts[name] for name in names if name != MAIN)
    document = Document(source)
    assert [p.text for p in document.paragraphs] == ["Alpha"]
    return source, parts, sealed_parts


def edit_alpha(staged):
    document = Document(staged)
    assert [p.text for p in document.paragraphs] == ["Alpha"]
    document.paragraphs[0].text = "Beta"
    document.save(staged)
    return {"success": True, "results": [{"target": MAIN, "success": True}]}


def test_canonical_preserve_unrelated(opc_case, tmp_path, request):
    case = opc_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._opc_preservation_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    source = destination = None
    source_bytes = None
    original_parts = sealed_parts = None
    result = None
    observed = None
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                source, original_parts, sealed_parts = authored_source(tmp_path)
                source_bytes = source.read_bytes()
                destination = tmp_path / "output.docx"
                case["sourceProfile"] = {"sealedFixtureId": FIXTURE_ID, "authoredMembers": 17,
                                          "opaqueBytes": len(original_parts[OPAQUE]),
                                          "sealedOpaqueSha256": THUMBNAIL_SHA256, "originalText": "Alpha"}
            elif index == 1:
                result = stage_patch(str(source), str(destination), "strict", 1, edit_alpha)
                assert result["success"] is True and result["changes_applied"] == 1, result
                assert result["changes_planned"] == 1 and destination.is_file()
                assert source.read_bytes() == source_bytes
                _, observed = archive_parts(destination)
                assert_graph(observed)
                assert [p.text for p in Document(destination).paragraphs] == ["Beta"]
                assert not list(tmp_path.glob(".office-patch-*"))
            elif index == 2:
                assert main_text(original_parts) == ["Alpha"]
                assert main_text(observed) == ["Beta"]
                assert original_parts[MAIN] != observed[MAIN]
                diff = result["package_diff"]
                assert MAIN in diff["changed"] and OPAQUE not in diff["changed"]
                assert diff["added"] == diff["removed"] == []
                assert set(observed) == set(original_parts)
                assert all(observed[name] == original_parts[name] for name in original_parts if name != MAIN)
                assert diff["changed_payload_hashes"][MAIN] == {
                    "before": hashlib.sha256(original_parts[MAIN]).hexdigest(),
                    "after": hashlib.sha256(observed[MAIN]).hexdigest()}
                case["observedEdit"] = {"sourceText": "Alpha", "savedText": "Beta",
                                        "changed": diff["changed"], "added": diff["added"], "removed": diff["removed"]}
            else:
                assert observed[OPAQUE] == original_parts[OPAQUE] == sealed_parts[OPAQUE]
                assert len(observed[OPAQUE]) == 8324
                assert observed["_rels/.rels"] == original_parts["_rels/.rels"]
                assert observed["[Content_Types].xml"] == original_parts["[Content_Types].xml"]
                assert source.read_bytes() == source_bytes
                case["observedCustody"] = {"opaqueBytes": len(observed[OPAQUE]),
                                           "opaqueIdentical": True, "sourceIdentical": True}
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_noop_cannot_satisfy_edited_text(tmp_path):
    source, parts, _ = authored_source(tmp_path)
    destination = tmp_path / "noop.docx"
    def noop(staged):
        assert Path(staged).read_bytes() == source.read_bytes()
        return {"success": True, "results": [{"target": MAIN, "success": True}]}
    result = stage_patch(str(source), str(destination), "strict", 1, noop)
    assert result["changes_applied"] == 1, result
    _, saved = archive_parts(destination)
    assert main_text(saved) == main_text(parts) == ["Alpha"]
    assert main_text(saved) != ["Beta"]
    assert saved[OPAQUE] == parts[OPAQUE]


def test_changed_opaque_payload_fails_byte_custody(tmp_path):
    source, parts, _ = authored_source(tmp_path)
    destination = tmp_path / "tampered.docx"
    def tamper(staged):
        result = edit_alpha(staged)
        staged_path = Path(staged)
        edited_members = []
        with zipfile.ZipFile(staged_path) as archive:
            for info in archive.infolist():
                payload = archive.read(info)
                if info.filename == OPAQUE:
                    payload = bytes([payload[0] ^ 1]) + payload[1:]
                edited_members.append((info, payload))
        with zipfile.ZipFile(staged_path, "w") as archive:
            for info, payload in edited_members:
                archive.writestr(info, payload)
        return result
    result = stage_patch(str(source), str(destination), "strict", 1, tamper)
    assert result["changes_applied"] == 1, result
    _, saved = archive_parts(destination)
    assert main_text(saved) == ["Beta"] and [p.text for p in Document(destination).paragraphs] == ["Beta"]
    assert len(saved[OPAQUE]) == len(parts[OPAQUE])
    assert saved[OPAQUE] != parts[OPAQUE]  # canonical byte assertion would refuse this edit
    assert OPAQUE in result["package_diff"]["changed"]


def test_failed_callback_keeps_source_and_destination(tmp_path):
    source, _, _ = authored_source(tmp_path)
    destination = tmp_path / "existing.docx"
    destination.write_bytes(source.read_bytes())
    before = {p.name: p.read_bytes() for p in tmp_path.iterdir()}
    def fail(staged):
        edit_alpha(staged)
        raise RuntimeError("injected failure after staged edit")
    result = stage_patch(str(source), str(destination), "strict", 1, fail)
    assert result["success"] is False and result["changes_applied"] == 0
    assert "injected failure after staged edit" in result["error"]
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == before
    assert not list(tmp_path.glob(".office-patch-*"))
