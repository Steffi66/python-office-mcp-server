"""Run one exact anchor-discovery case with source-order and negative controls."""

from docx import Document

from tests.docx_anchor_headings_paragraphs.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools

CONTENT = ("Introduction", "Customer context paragraph", "Delivery approach", "Use iterative delivery")


def _saved(path):
    doc = Document()
    doc.add_heading(CONTENT[0], level=1)
    doc.add_paragraph(CONTENT[1])
    doc.add_heading(CONTENT[2], level=1)
    doc.add_paragraph(CONTENT[3])
    doc.save(path)
    assert path.is_file()
    assert tuple(p.text for p in Document(path).paragraphs if p.text) == CONTENT


def _anchors(api, path, query=None):
    before = path.read_bytes()
    response = api.tool_word_list_anchors(str(path), query=query)
    assert path.read_bytes() == before
    assert "error" not in response and response["query"] == query
    assert response["count"] == len(response["anchors"])
    return response


def test_canonical_headings_and_paragraphs(anchor_headings_paragraphs_case, tmp_path, request):
    case = anchor_headings_paragraphs_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_anchor_headings_paragraphs_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "anchors.docx"
    api = WordAdvancedTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                _saved(path)
                source = path.read_bytes()
            elif index == 1:
                response = _anchors(api, path)
                assert response["file"] == str(path)
            elif index == 2:
                assert response["count"] >= 4 and len(response["anchors"]) >= 4
                anchors = response["anchors"]
                assert [entry["anchor_text"] for entry in anchors[:4]] == list(CONTENT)
                indexes = [entry["paragraph_index"] for entry in anchors[:4]]
                assert all(isinstance(value, int) for value in indexes)
                assert indexes == sorted(indexes) and len(set(indexes)) == 4
            elif index == 3:
                assert any(entry["type"] == "section_heading" and entry["anchor_text"] == "Introduction"
                           for entry in anchors)
                assert anchors[0]["type"] == "section_heading"
            else:
                assert any(entry["type"] == "paragraph" and "Customer context" in entry["anchor_text"]
                           for entry in anchors)
                assert anchors[1]["type"] == "paragraph" and anchors[1]["anchor_text"] == CONTENT[1]
                assert path.read_bytes() == source
                case["observedAnchors"] = [entry["anchor_text"] for entry in anchors[:4]]
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_query_and_paragraph_controls_are_not_default_claim(tmp_path):
    path = tmp_path / "query-controls.docx"
    _saved(path)
    api = WordAdvancedTools()
    default = _anchors(api, path)
    assert [entry["type"] for entry in default["anchors"][:4]] == [
        "section_heading", "paragraph", "section_heading", "paragraph"]
    headings = _anchors(api, path, query="delivery")
    assert headings["count"] == 2
    assert [entry["anchor_text"] for entry in headings["anchors"]] == [CONTENT[2], CONTENT[3]]
    assert all("delivery" in entry["anchor_text"].casefold() for entry in headings["anchors"])
    no_paragraphs = api.tool_word_list_anchors(str(path), include_paragraphs=False)
    assert no_paragraphs["count"] == 2
    assert [entry["anchor_text"] for entry in no_paragraphs["anchors"]] == [CONTENT[0], CONTENT[2]]
    assert _anchors(api, path, query="unmatched token")["count"] == 0
