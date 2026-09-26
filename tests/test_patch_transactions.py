"""Saved-document outcomes for staged batches, independent of acceptance bindings."""

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from office_server import OfficeServer
from tools.word_advanced_tools import _get_text_with_track_changes


def make_document(path):
    if path.suffix == ".xlsx":
        doc = Workbook()
        doc.active["A1"] = "before"
    elif path.suffix == ".docx":
        doc = Document()
        doc.add_paragraph("<Present>")
    else:
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[0])
        slide.shapes.title.text = "Original title"
        slide.placeholders[1].text = "Original subtitle"
    doc.save(path)


def snapshot(directory):
    return {p.name: p.read_bytes() for p in directory.iterdir() if p.is_file()}


def destination_for(source, destination):
    if destination == "source":
        return source
    output = source.with_name("output" + source.suffix)
    if destination == "existing":
        output.write_bytes(b"existing output must survive failure")
    return output


@pytest.mark.parametrize("suffix,target", [(".xlsx", "A1"), (".docx", "<Present>"), (".pptx", "slide:1/title")])
@pytest.mark.parametrize("destination", ["source", "absent", "existing"])
def test_preview_never_modifies_source_or_destination(tmp_path, suffix, target, destination):
    source = tmp_path / ("input" + suffix)
    make_document(source)
    output = destination_for(source, destination)
    before = snapshot(tmp_path)
    entries = sorted(p.name for p in tmp_path.iterdir())
    result = OfficeServer().tool_office_patch(
        str(source), [{"target": target, "value": "changed"}],
        mode="dry_run", output_path=str(output),
    )
    assert result["success"], result
    assert result["changes_applied"] == 0
    assert result["changes_planned"] == 1
    assert not result["results"][0]["applied"]
    assert snapshot(tmp_path) == before
    assert sorted(p.name for p in tmp_path.iterdir()) == entries


@pytest.mark.parametrize("destination", ["source", "absent", "existing"])
@pytest.mark.parametrize("suffix,good,bad", [
    (".xlsx", "A1", "Missing!B1"),
    (".docx", "<Present>", "<Missing>"),
    (".pptx", "slide:1/title", "slide:1/absent shape"),
])
def test_strict_refusal_is_atomic(tmp_path, destination, suffix, good, bad):
    source = tmp_path / ("input" + suffix)
    make_document(source)
    output = destination_for(source, destination)
    before = snapshot(tmp_path)
    result = OfficeServer().tool_office_patch(
        str(source), [{"target": good, "value": "changed"}, {"target": bad, "value": "missing"}],
        mode="strict", output_path=str(output),
    )
    assert result["success"] is False, result
    assert result["changes_applied"] == 0
    assert all(not r["applied"] for r in result.get("results", []))
    assert snapshot(tmp_path) == before
    assert not any(p.is_dir() for p in tmp_path.iterdir())


@pytest.mark.parametrize("destination", ["absent", "existing"])
def test_presentation_batch_accumulates_at_distinct_output(tmp_path, destination):
    source = tmp_path / "input.pptx"
    make_document(source)
    original = source.read_bytes()
    output = destination_for(source, destination)
    result = OfficeServer().tool_office_patch(
        str(source), [
            {"target": "slide:1/title", "value": "Changed title"},
            {"target": "slide:1/subtitle", "value": "Changed subtitle"},
        ], output_path=str(output), mode="safe",
    )
    assert result["changes_applied"] == 2, result
    slide = Presentation(output).slides[0]
    assert slide.shapes.title.text == "Changed title"
    assert slide.placeholders[1].text == "Changed subtitle"
    assert source.read_bytes() == original


def test_word_best_effort_counts_each_placeholder(tmp_path):
    source = tmp_path / "input.docx"
    make_document(source)
    result = OfficeServer().tool_office_patch(str(source), [
        {"target": "<Present>", "value": "changed"},
        {"target": "<Missing>", "value": "absent"},
    ])
    assert result["status"] == "partial_success", result
    assert result["changes_applied"] == 1
    assert [r["applied"] for r in result["results"]] == [True, False]
    assert "".join(_get_text_with_track_changes(p) for p in Document(source).paragraphs) == "changed"


@pytest.mark.parametrize("value", [[["bad", "partial"], ["short"]], [["bad", "partial"], None], [["bad", "partial"], [1, {}]]])
def test_invalid_range_does_not_leak_partial_rows_in_best_effort(tmp_path, value):
    source = tmp_path / "input.xlsx"
    make_document(source)
    result = OfficeServer().tool_office_patch(str(source), [
        {"target": "A1:B2", "value": value},
        {"target": "D1", "value": "valid"},
    ])
    assert result["changes_applied"] == 1, result
    workbook = load_workbook(source)
    try:
        assert workbook.active["A1"].value == "before"
        assert workbook.active["B1"].value is None
        assert workbook.active["D1"].value == "valid"
    finally:
        workbook.close()


def test_missing_pptx_placeholder_is_not_applied(tmp_path):
    source = tmp_path / "input.pptx"
    make_document(source)
    before = source.read_bytes()
    result = OfficeServer().tool_office_patch(str(source), [{"target": "absent", "value": "changed"}])
    assert result["changes_applied"] == 0
    assert source.read_bytes() == before


@pytest.mark.parametrize("bad_target", ["A0", "XFE1", "A1junk", "A2:A1", "A0:B1"])
def test_invalid_excel_address_preserves_source(tmp_path, bad_target):
    source = tmp_path / "input.xlsx"
    make_document(source)
    before = source.read_bytes()
    value = [[1, 2]] if ":" in bad_target else "new"
    result = OfficeServer().tool_office_patch(str(source), [{"target": bad_target, "value": value}])
    assert result["changes_applied"] == 0, result
    assert source.read_bytes() == before
