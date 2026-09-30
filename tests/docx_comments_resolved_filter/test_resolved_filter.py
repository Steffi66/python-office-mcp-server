"""Execute one authored resolved-filter case with separate filter/refusal controls."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_resolved_filter.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _comments(response, expected_filter, count):
    assert "error" not in response and response.get("filter") == expected_filter
    records = response["comments"]
    assert response["comment_count"] == len(records) == count
    return records


def _author_two(path, authoring):
    doc = Document()
    doc.add_paragraph("Alpha target")
    doc.add_paragraph("Beta target")
    doc.save(path)
    assert path.is_file() and [p.text for p in Document(path).paragraphs if p.text] == ["Alpha target", "Beta target"]
    alpha = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Alpha target", comment_text="Comment alpha", author="Manuel")
    beta = authoring.tool_word_add_comment(
        file_path=str(path), target_text="Beta target", comment_text="Comment beta", author="Rui Carmo")
    assert alpha.get("success") is True and beta.get("success") is True


def _returned_ids_and_anchors(path, reader):
    before = _comments(reader.tool_word_get_comments(str(path), filter="all"), "all", 2)
    assert [(r["text"], r["author"], r["done"]) for r in before] == [
        ("Comment alpha", "Manuel", False), ("Comment beta", "Rui Carmo", False)]
    first_id, second_id = (str(r["id"]) for r in before)
    assert first_id != second_id
    body = etree.fromstring(_parts(path)["word/document.xml"])
    paragraphs = body.findall(".//" + W + "body/" + W + "p")
    targets = [p for p in paragraphs if "".join(p.itertext()) in ("Alpha target", "Beta target")]
    assert ["".join(p.itertext()) for p in targets] == ["Alpha target", "Beta target"]
    for target, expected_id in zip(targets, (first_id, second_id)):
        assert [n.get(W + "id") for n in target.findall(".//" + W + "commentRangeStart")] == [expected_id]
        assert [n.get(W + "id") for n in target.findall(".//" + W + "commentRangeEnd")] == [expected_id]
    return first_id, second_id


def test_canonical_authored_resolved_filter(resolved_filter_case, tmp_path, request):
    case = resolved_filter_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_resolved_filter_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "resolved-filter.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                _author_two(path, authoring)
            elif index == 1:
                first_id, second_id = _returned_ids_and_anchors(path, reader)
                case["selectedId"] = first_id
                original_parts = _parts(path)
            elif index == 2:
                response = reader.tool_word_resolve_comment(str(path), first_id, True)
                assert path.is_file()
            elif index == 3:
                assert response.get("success") is True and response.get("done") is True
                assert str(response.get("thread_root_comment_id")) == first_id
            else:
                # The production filter is exercised on a fresh package read.
                resolved = _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 1)
                assert [(str(r["id"]), r["text"], r["author"], r["done"]) for r in resolved] == [
                    (first_id, "Comment alpha", "Manuel", True)]
                unfiltered = _comments(reader.tool_word_get_comments(str(path), filter="all"), "all", 2)
                assert [(str(r["id"]), r["done"]) for r in unfiltered] == [(first_id, True), (second_id, False)]
                changed_parts = _parts(path)
                old_comments = etree.fromstring(original_parts["word/comments.xml"])
                new_comments = etree.fromstring(changed_parts["word/comments.xml"])
                assert [(n.get(W + "id"), "".join(n.itertext())) for n in old_comments] == [
                    (n.get(W + "id"), "".join(n.itertext())) for n in new_comments]
                extension = etree.fromstring(changed_parts["word/commentsExtended.xml"])
                assert any(node.get(W15 + "paraId") == str(unfiltered[0]["para_id"])
                           and node.get(W15 + "done") == "1" for node in extension.findall(W15 + "commentEx"))
                case["observedResolvedId"] = first_id
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_filter_excludes_open_and_opposite_resolution_controls(tmp_path):
    path = tmp_path / "opposite.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    _author_two(path, authoring)
    first_id, second_id = _returned_ids_and_anchors(path, reader)
    assert _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 0) == []
    response = reader.tool_word_resolve_comment(str(path), second_id, True)
    assert response.get("success") is True and str(response.get("thread_root_comment_id")) == second_id
    resolved = _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 1)
    assert [(str(r["id"]), r["done"]) for r in resolved] == [(second_id, True)]
    open_rows = _comments(reader.tool_word_get_comments(str(path), filter="open"), "open", 1)
    assert [(str(r["id"]), r["done"]) for r in open_rows] == [(first_id, False)]
    before_noop = path.read_bytes()
    again = reader.tool_word_resolve_comment(str(path), second_id, True)
    assert again.get("success") is True and again.get("unchanged") is True
    assert path.read_bytes() == before_noop


def test_missing_id_does_not_publish(tmp_path):
    path = tmp_path / "missing.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    _author_two(path, authoring)
    first_id, second_id = _returned_ids_and_anchors(path, reader)
    source = path.read_bytes()
    unknown = str(max(int(first_id), int(second_id)) + 100)
    response = reader.tool_word_resolve_comment(str(path), unknown, True)
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source
    assert _comments(reader.tool_word_get_comments(str(path), filter="resolved"), "resolved", 0) == []
