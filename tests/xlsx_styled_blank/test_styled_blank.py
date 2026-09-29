"""One authored styled blank cell; independent ZIP/XML and mutation controls."""

import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from tests.xlsx_styled_blank.cases import CASE_KEY
from tools.excel_advanced_tools import ExcelAdvancedTools

SHEET = "xl/worksheets/sheet1.xml"
STYLE = "xl/styles.xml"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def make_source(path, *, styled=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Styled"
    cell = sheet["A1"]
    if styled:
        cell.fill = PatternFill(fill_type="solid", fgColor="FF1177CC")
        cell.font = Font(bold=True, color="FF224466")
        cell.number_format = "0.00"
        cell.alignment = Alignment(horizontal="center")
    sheet["B1"] = "guard"
    workbook.save(path)
    workbook.close()


def cell_facts(path):
    workbook = load_workbook(path)
    try:
        sheet = workbook["Styled"]
        cell = sheet["A1"]
        return (cell.value, cell.has_style, cell.style_id, cell.fill.fgColor.rgb,
                cell.font.bold, cell.font.color.rgb if cell.font.color.type == "rgb" else None,
                cell.number_format, cell.alignment.horizontal, sheet["B1"].value)
    finally:
        workbook.close()


def archive_facts(path):
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names)) == 9
        parts = {name: archive.read(name) for name in names}
    root = etree.fromstring(parts[SHEET])
    cells = root.xpath("//m:sheetData/m:row/m:c[@r='A1']", namespaces=NS)
    assert len(cells) == 1
    return names, parts, cells[0]


def test_canonical_styled_blank(styled_case, tmp_path, request):
    case = styled_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._xlsx_styled_blank_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    source = tmp_path / "source.xlsx"
    output = tmp_path / "published.xlsx"
    original_bytes = original_names = original_parts = original_cell = original_facts = None
    read_result = patch_result = saved_facts = saved_names = saved_parts = saved_cell = None
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                make_source(source)  # Synthetic temp input; no prebuilt shared styled-blank fixture.
                original_bytes = source.read_bytes()
                original_names, original_parts, original_cell = archive_facts(source)
                original_facts = cell_facts(source)
                assert original_facts == (None, True, 1, "FF1177CC", True, "FF224466", "0.00", "center", "guard")
                assert original_cell.get("s") == "1" and original_cell.get("t") == "n"
                assert not original_cell.xpath("./m:is|./m:v", namespaces=NS)
                assert not output.exists()
                case["sourceProfile"] = {"synthetic": True, "parts": len(original_names), "A1": None,
                                         "styleRef": "1", "B1": "guard"}
            elif index == 1:
                tools = ExcelAdvancedTools()
                read_result = tools.tool_excel_get_range(str(source), "A1", sheet_name="Styled")
                patch_result = tools.tool_excel_patch_cell(
                    str(source), "A1", "filled", sheet_name="Styled", log_change=False,
                    highlight=False, output_path=str(output))
                assert output.is_file() and source.read_bytes() == original_bytes
                saved_facts = cell_facts(output)  # Independent reopen, not the writer's result.
                saved_names, saved_parts, saved_cell = archive_facts(output)
            elif index == 2:
                assert read_result == {"sheet": "Styled", "range": "A1:A1", "row_count": 1,
                                       "col_count": 1, "data": [[{"ref": "A1", "value": None, "type": "empty"}]]}
                assert patch_result["success"] is True and patch_result["sheet"] == "Styled"
                assert patch_result["cell"] == "A1" and patch_result["old_value"] is None
                case["observedRead"] = {"A1": read_result["data"][0][0], "oldValue": patch_result["old_value"]}
            else:
                assert patch_result["new_value"] == "filled" and patch_result["value_type"] == "str"
                assert patch_result["logged"] is False and patch_result["highlighted"] is False
                assert patch_result["file"] == str(output)
                assert saved_facts == ("filled", *original_facts[1:])
                assert saved_cell.get("s") == original_cell.get("s") == "1"
                assert saved_cell.get("t") == "inlineStr"
                assert saved_cell.xpath("string(./m:is/m:t)", namespaces=NS) == "filled"
                assert saved_names == original_names and set(saved_parts) == set(original_parts)
                assert saved_parts[STYLE] == original_parts[STYLE]
                assert {name for name in original_names if saved_parts[name] != original_parts[name]} == {SHEET}
                assert output.read_bytes() != original_bytes and source.read_bytes() == original_bytes
                assert archive_facts(source)[1] == original_parts
                assert cell_facts(source) == original_facts
                case["observedSave"] = {"value": saved_facts[0], "styleRef": saved_cell.get("s"),
                                        "styleComponentsRetained": True, "changedParts": [SHEET],
                                        "stylesPartIdentical": True, "sourceBytesIdentical": True,
                                        "B1": saved_facts[-1]}
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_unstyled_wrong_target_highlight_and_wrong_value_do_not_satisfy_case(tmp_path):
    tools = ExcelAdvancedTools()
    source = tmp_path / "styled.xlsx"
    plain = tmp_path / "plain.xlsx"
    make_source(source)
    make_source(plain, styled=False)
    original_bytes = source.read_bytes()
    assert cell_facts(plain)[1] is False  # Null alone cannot establish styled blank.

    wrong_target = tmp_path / "wrong-target.xlsx"
    assert tools.tool_excel_patch_cell(str(source), "B1", "filled", sheet_name="Styled",
                                       log_change=False, highlight=False, output_path=str(wrong_target))["success"]
    assert cell_facts(wrong_target)[0] is None  # A1 still blank.

    highlighted = tmp_path / "highlighted.xlsx"
    assert tools.tool_excel_patch_cell(str(source), "A1", "filled", sheet_name="Styled",
                                       log_change=False, highlight=True, output_path=str(highlighted))["success"]
    assert cell_facts(highlighted)[3] != cell_facts(source)[3]  # Original blue fill lost.

    wrong_value = tmp_path / "wrong-value.xlsx"
    assert tools.tool_excel_patch_cell(str(source), "A1", "different", sheet_name="Styled",
                                       log_change=False, highlight=False, output_path=str(wrong_value))["success"]
    assert cell_facts(wrong_value)[0] == "different" != "filled"
    assert source.read_bytes() == original_bytes
