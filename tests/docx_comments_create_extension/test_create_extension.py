"""Run the exact missing commentsExtended creation case with independent controls."""

from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.docx_comments_create_extension.cases import CASE_KEY
from tools.word_advanced_tools import WordAdvancedTools
from tools.word_tools import WordTools

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
EXTENSION = "word/commentsExtended.xml"
REL_TYPE = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
CT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"


def _parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        assert names and len(names) == len(set(names))
        return {name: archive.read(name) for name in names}


def _read(reader, path):
    result = reader.tool_word_get_comments(str(path), filter="all")
    assert "error" not in result and result.get("filter") == "all"
    assert result.get("comment_count") == len(result["comments"])
    return result["comments"]


def _create(path, authoring, reader, *, second=False):
    document = Document()
    document.add_paragraph("No commentsExtended yet")
    if second:
        document.add_paragraph("Another target")
    document.save(path)
    assert [p.text for p in Document(path).paragraphs if p.text] == (
        ["No commentsExtended yet", "Another target"] if second else ["No commentsExtended yet"])
    result = authoring.tool_word_add_comment(
        file_path=str(path), target_text="No commentsExtended yet", comment_text="Needs follow-up")
    assert result.get("success") is True
    if second:
        other = authoring.tool_word_add_comment(
            file_path=str(path), target_text="Another target", comment_text="Other comment")
        assert other.get("success") is True
    rows = _read(reader, path)
    assert len(rows) == (2 if second else 1)
    assert rows[0]["text"] == "Needs follow-up" and rows[0]["done"] is False
    if second:
        assert rows[1]["text"] == "Other comment" and rows[1]["done"] is False
    before = _parts(path)
    assert EXTENSION not in before
    return str(rows[0]["id"]), rows, before


def _extension_entry(parts, para_id):
    assert EXTENSION in parts
    root = etree.fromstring(parts[EXTENSION])
    assert root.tag == W15 + "commentsEx"
    entries = root.findall(W15 + "commentEx")
    matching = [entry for entry in entries if entry.get(W15 + "paraId") == para_id]
    assert len(matching) == 1
    return entries, matching[0]


def _assert_package_metadata(parts):
    types = etree.fromstring(parts["[Content_Types].xml"])
    overrides = [node for node in types.findall(CT + "Override")
                 if node.get("PartName") == "/" + EXTENSION]
    assert len(overrides) == 1 and overrides[0].get("ContentType") == CT_TYPE
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    links = [node for node in rels.findall(REL + "Relationship") if node.get("Type") == REL_TYPE]
    assert len(links) == 1 and links[0].get("Target") == "commentsExtended.xml" and links[0].get("Id")


def test_canonical_create_extension(create_extension_case, tmp_path, request):
    case = create_extension_case
    assert case["stableCaseKey"] == CASE_KEY
    ledger = request.config._docx_comments_create_extension_ledger
    case.update(outcome="running", nodeid=request.node.nodeid)
    ledger.write()
    path = tmp_path / "new-extension.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    for index, step in enumerate(case["steps"]):
        try:
            if index == 0:
                comment_id, original, before = _create(path, authoring, reader)
                assert len(original) == 1 and str(original[0]["id"]) == comment_id
                case["commentId"] = comment_id
            elif index == 1:
                assert EXTENSION not in before
                assert EXTENSION not in _parts(path)
                assert original[0]["done"] is False
            elif index == 2:
                resolved = reader.tool_word_resolve_comment(str(path), comment_id, True)
                assert path.is_file()
            elif index == 3:
                assert resolved.get("success") is True and resolved.get("unchanged") is not True
                assert str(resolved.get("comment_id")) == comment_id
                assert str(resolved.get("thread_root_comment_id")) == comment_id
                assert resolved.get("done") is True
            elif index == 4:
                fresh = _read(reader, path)
                assert len(fresh) == 1 and str(fresh[0]["id"]) == comment_id
                assert fresh[0]["done"] is True
                para_id = fresh[0]["para_id"]
                assert isinstance(para_id, str) and para_id
                comments = etree.fromstring(_parts(path)["word/comments.xml"])
                node = [node for node in comments.findall(".//" + W + "comment")
                        if node.get(W + "id") == comment_id]
                assert len(node) == 1 and node[0].find(W + "p") is not None
                assert node[0].find(W + "p").get(W14 + "paraId") == para_id
                case["paraId"] = para_id
            else:
                after = _parts(path)
                entries, entry = _extension_entry(after, para_id)
                assert len(entries) == 1 and entry.get(W15 + "done") == "1"
                _assert_package_metadata(after)
                assert after["word/document.xml"] == before["word/document.xml"]
                case["extensionCreated"] = True
            step["outcome"] = "passed"
        except BaseException as exc:
            step.update(outcome="failed", error=str(exc))
            case["outcome"] = "failed"
            for following in case["steps"]:
                if following["outcome"] == "not-run":
                    following["outcome"] = "skipped"
            ledger.write()
            raise
        ledger.write()
    case["outcome"] = "passed"
    ledger.write()


def test_second_comment_is_not_marked_done_by_new_extension(tmp_path):
    path = tmp_path / "two-comments.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    first, original, before = _create(path, authoring, reader, second=True)
    second = str(original[1]["id"])
    assert first != second
    response = reader.tool_word_resolve_comment(str(path), first, True)
    assert response.get("success") is True
    rows = _read(reader, path)
    assert [(str(row["id"]), row["done"]) for row in rows] == [(first, True), (second, False)]
    assert rows[0]["para_id"] and rows[0]["para_id"] != rows[1]["para_id"]
    entries, selected = _extension_entry(_parts(path), rows[0]["para_id"])
    assert len(entries) == 1 and selected.get(W15 + "done") == "1"
    assert entries[0].get(W15 + "paraId") != rows[1]["para_id"]
    assert _parts(path)["word/document.xml"] == before["word/document.xml"]


def test_unknown_id_does_not_create_extension(tmp_path):
    path = tmp_path / "unknown.docx"
    authoring, reader = WordAdvancedTools(), WordTools()
    comment_id, _, _ = _create(path, authoring, reader)
    source = path.read_bytes()
    unknown = str(int(comment_id) + 1000)
    response = reader.tool_word_resolve_comment(str(path), unknown, True)
    assert "error" in response and response.get("success") is not True
    assert path.read_bytes() == source and EXTENSION not in _parts(path)
    assert _read(reader, path)[0]["done"] is False
