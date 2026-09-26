"""Saved-document text and clone outcomes on the native Python APIs."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from office_server import OfficeServer
from tools.word_advanced_tools import _get_text_with_track_changes


def test_word_split_span_preserves_boundary_fonts_and_revisions(tmp_path):
    path = tmp_path / "word.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("prefix <Cus").bold = True
    p.add_run("tomer> suffix").italic = True
    doc.save(path)
    result = OfficeServer().tool_office_patch(str(path), [{"target": "<Customer>", "value": "Acme"}])
    assert result["changes_applied"] == 1, result
    p = next(p for p in Document(path).paragraphs if _get_text_with_track_changes(p))
    assert _get_text_with_track_changes(p) == "prefix Acme suffix"
    assert p.runs[0].text == "prefix " and p.runs[0].bold
    assert p.runs[-1].text == " suffix" and p.runs[-1].italic
    deleted = p._p.findall(qn("w:del"))
    assert "".join(n.text or "" for n in deleted[0].iter(qn("w:delText"))) == "<Customer>"
    assert len(deleted[0].findall(qn("w:r"))) == 2
    OfficeServer().tool_word_accept_all_changes(str(path))
    assert "prefix Acme suffix" in [p.text for p in Document(path).paragraphs]


def test_word_span_cannot_cross_field_barrier(tmp_path):
    path = tmp_path / "word.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("<Cus")
    p._p.append(OxmlElement("w:fldSimple"))
    p.add_run("tomer>")
    doc.save(path)
    before = path.read_bytes()
    result = OfficeServer().tool_office_patch(str(path), [{"target": "<Customer>", "value": "Acme"}], mode="strict")
    assert result["changes_applied"] == 0
    assert path.read_bytes() == before


def test_slide_split_run_replacement_preserves_properties(tmp_path):
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    p = prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text_frame.paragraphs[0]
    p.clear()
    a, b = p.add_run(), p.add_run()
    a.text, b.text = "pre <Cus", "tomer> post <Customer>"
    a.font.bold, b.font.italic = True, True
    prs.save(path)
    result = OfficeServer().tool_office_patch(str(path), [{"target": "<Customer>", "value": "Acme"}])
    assert result["changes_applied"] == 1, result
    p = Presentation(path).slides[0].shapes.title.text_frame.paragraphs[0]
    assert p.text == "pre Acme post Acme"
    assert p.runs[0].font.bold and p.runs[-1].font.italic


def test_duplicate_chart_has_independent_workbook_and_chart_part(tmp_path):
    path = tmp_path / "chart.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Series", [1, 2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data)
    prs.save(path)
    result = OfficeServer().tool_pptx_duplicate_slide(str(path), 1)
    assert result["success"], result
    prs = Presentation(path)
    original, clone = prs.slides[0].shapes[0].chart, prs.slides[1].shapes[0].chart
    assert original.part.partname != clone.part.partname
    assert original.part.chart_workbook.xlsx_part.partname != clone.part.chart_workbook.xlsx_part.partname
    changed = CategoryChartData()
    changed.categories = ["A", "B"]
    changed.add_series("Series", [9, 8])
    clone.replace_data(changed)
    prs.save(path)
    prs = Presentation(path)
    assert list(prs.slides[0].shapes[0].chart.series[0].values) == [1, 2]
    assert list(prs.slides[1].shapes[0].chart.series[0].values) == [9, 8]


def test_import_existing_notes_is_explicit_and_donor_unchanged(tmp_path):
    source, target = tmp_path / "donor.pptx", tmp_path / "target.pptx"
    donor = Presentation()
    slide = donor.slides.add_slide(donor.slide_layouts[0])
    slide.shapes.title.text = "Donor"
    slide.notes_slide.notes_text_frame.text = "Private note"
    donor.save(source)
    receiver = Presentation()
    receiver.slides.add_slide(receiver.slide_layouts[0])
    receiver.save(target)
    before = source.read_bytes()
    result = OfficeServer().tool_pptx_import_slide(str(source), 1, str(target), include_notes=True)
    assert result["success"], result
    assert source.read_bytes() == before
    imported = Presentation(target).slides[-1]
    assert "Private note" in imported.notes_slide.notes_text_frame.text
