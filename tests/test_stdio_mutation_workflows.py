"""Real MCP stdio workflow with bounded I/O; optionally target a clean wheel install."""

import json
import os
import queue
import subprocess
import sys
import threading
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation


class Client:
    def __init__(self, command, cwd):
        self.process = subprocess.Popen(command, cwd=cwd, stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                        stderr=subprocess.PIPE, text=True)
        self.responses = queue.Queue()
        self.errors = []
        self.next_id = 0
        self.transcript = []
        self.readers = [threading.Thread(target=self._read, daemon=True),
                        threading.Thread(target=self._errors, daemon=True)]
        for thread in self.readers:
            thread.start()

    def _read(self):
        for line in self.process.stdout:
            self.responses.put(line)

    def _errors(self):
        for line in self.process.stderr:
            self.errors.append(line)
            self.errors[:] = self.errors[-100:]

    def request(self, method, params):
        self.next_id += 1
        request = {"jsonrpc": "2.0", "id": self.next_id, "method": method, "params": params}
        self.process.stdin.write(json.dumps(request) + "\n")
        self.process.stdin.flush()
        try:
            response = json.loads(self.responses.get(timeout=12))
        except queue.Empty as exc:
            raise AssertionError("MCP response timeout: " + "".join(self.errors)) from exc
        self.transcript.append({"request": request, "response": response})
        assert response["id"] == request["id"]
        assert "error" not in response, response
        return response["result"]

    def call(self, name, **args):
        response = self.request("tools/call", {"name": name, "arguments": args})
        decoded = json.loads(response["content"][0]["text"])
        if isinstance(decoded, dict):
            assert response["structuredContent"] == decoded
            assert response["isError"] == bool(decoded.get("error") or decoded.get("success") is False)
        return decoded

    def close(self):
        self.process.terminate()
        try:
            self.process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            self.process.kill()
            self.process.wait(timeout=3)
        for stream in (self.process.stdin, self.process.stdout, self.process.stderr):
            stream.close()
        for thread in self.readers:
            thread.join(timeout=1)


@pytest.fixture
def client(tmp_path):
    python = os.environ.get("OFFICE_MCP_TEST_PYTHON", sys.executable)
    command = [python, "-m", "office_server"] if "OFFICE_MCP_TEST_PYTHON" in os.environ else [python, str(Path(__file__).parents[1] / "office_server.py")]
    client = Client(command, tmp_path)
    try:
        init = client.request("initialize", {"protocolVersion": "2024-11-05", "capabilities": {}, "clientInfo": {"name": "acceptance", "version": "1"}})
        assert init["serverInfo"]
        client.process.stdin.write(json.dumps({"jsonrpc": "2.0", "method": "notifications/initialized"}) + "\n")
        client.process.stdin.flush()
        tools = client.request("tools/list", {})
        assert {"office_inspect", "office_read", "office_patch"}.issubset({t["name"] for t in tools["tools"]})
        yield client
    finally:
        target = Path(__file__).parents[1] / "test-results" / "stdio"
        target.mkdir(parents=True, exist_ok=True)
        (target / (tmp_path.name + ".json")).write_text(json.dumps(client.transcript, indent=2))
        client.close()


@pytest.mark.parametrize("suffix,target,value,inspect", [
    (".xlsx", "A1", "first\nsecond", "sheets"),
    (".docx", "<Present>", "Changed", "sections"),
    (".pptx", "slide:1/title", "Changed", "slides"),
])
def test_inspect_preview_patch_read_over_stdio(client, tmp_path, suffix, target, value, inspect):
    source = tmp_path / ("input" + suffix)
    output = tmp_path / ("output" + suffix)
    if suffix == ".xlsx":
        document = Workbook()
        document.active["A1"] = "before"
    elif suffix == ".docx":
        document = Document()
        document.add_paragraph("<Present>")
    else:
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = "before"
    document.save(source)
    before = source.read_bytes()
    assert "error" not in client.call("office_inspect", file_path=str(source), what=inspect)
    args = {"file_path": str(source), "output_path": str(output), "changes": [{"target": target, "value": value}]}
    preview = client.call("office_patch", **args, mode="dry_run")
    assert preview["success"] and preview["changes_applied"] == 0, preview
    assert source.read_bytes() == before and not output.exists()
    result = client.call("office_patch", **args, mode="safe")
    assert result["changes_applied"] == 1, result
    assert source.read_bytes() == before
    read = client.call("office_read", file_path=str(output))
    assert "error" not in read, read
    assert value in json.dumps(read, ensure_ascii=False).replace("\\n", "\n"), read
    if suffix == ".xlsx":
        wb = load_workbook(output)
        try:
            assert wb.active["A1"].value == value
            assert wb.active["A1"].alignment.wrap_text
        finally:
            wb.close()


def test_strict_refusal_over_stdio_preserves_existing_files(client, tmp_path):
    source, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    wb = Workbook()
    wb.active["A1"] = "before"
    wb.save(source)
    output.write_bytes(b"original destination")
    before = source.read_bytes()
    result = client.call("office_patch", file_path=str(source), output_path=str(output), mode="strict", changes=[
        {"target": "A1", "value": "changed"}, {"target": "Missing!A1", "value": 1},
    ])
    assert result["success"] is False and result["changes_applied"] == 0
    assert source.read_bytes() == before
    assert output.read_bytes() == b"original destination"
