"""Private staging for unified mutations; only a validated batch reaches its destination."""

from __future__ import annotations

import hashlib
import inspect
import os
import shutil
import tempfile
import threading
import zipfile
from collections.abc import Callable
from contextlib import contextmanager
from contextvars import ContextVar
from functools import wraps
from pathlib import Path
from typing import Any

from umcp_shared import MCPRequestCancelled, raise_if_cancelled


def validate_staged_document(path: Path) -> None:
    """Check archive readability and reopen through the format library before publishing."""
    from .package_guard import admit_package

    admit_package(path)
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError("duplicate package members")
        if archive.testzip() is not None:
            raise ValueError("invalid package member checksum")
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        Document(path)
    elif suffix == ".pptx":
        from pptx import Presentation

        Presentation(path)
    else:
        from openpyxl import load_workbook

        from tools.excel_advanced_tools import _close_workbook

        workbook = load_workbook(path, keep_vba=suffix in {".xlsm", ".xltm"})
        _close_workbook(workbook)


_staged_paths: ContextVar[frozenset[str]] = ContextVar("office_staged_paths", default=frozenset())


def staged_writer(function=None, *, source_argument="file_path", read_operations=()):
    """Preserve tool signatures while enrolling an explicit writer in private staging.

    Only a nested call addressing the same private document bypasses restaging.
    Public methods remain usable directly as well as through the MCP dispatcher.
    """
    def decorate(method):
        signature = inspect.signature(method)

        @wraps(method)
        def wrapped(self, *args, **kwargs):
            bound = signature.bind(self, *args, **kwargs)
            bound.apply_defaults()
            values = dict(bound.arguments)
            values.pop("self")
            if values.get("operation") in read_operations:
                return method(self, **values)
            from .save_utils import resolve_office_path

            source = str(Path(resolve_office_path(values[source_argument])).resolve())
            if source in _staged_paths.get():
                return method(self, **values)
            mode = values.get("mode", "best_effort")
            if mode not in {"best_effort", "safe", "strict", "dry_run"}:
                return {"success": False, "error": "Unsupported mutation mode", "changes_applied": 0}
            data = values.get("data")
            nested_output = data.get("output_path") if isinstance(data, dict) else None
            output = values.get("output_path") or nested_output
            if output:
                output = str(Path(resolve_office_path(output)).resolve())
            if mode == "safe" and (not output or Path(output).resolve() == Path(source)):
                return {"success": False, "mode": mode, "status": "failed", "error": "safe mode requires a distinct output_path", "changes_applied": 0}

            def apply(staged):
                call = dict(values)
                call[source_argument] = staged
                call["output_path"] = None
                if "mode" in call:
                    call["mode"] = "strict" if mode == "strict" else "best_effort"
                if nested_output:
                    call["data"] = {key: value for key, value in data.items() if key != "output_path"}
                before = Path(staged).read_bytes()
                result = method(self, **call)
                if not isinstance(result, dict):
                    raise ValueError("Writer must return structured diagnostics")
                if (method.__name__ == "tool_word_resolve_comment" and result.get("unchanged")
                        and output and Path(output).resolve() != Path(source)):
                    # A no-op cannot name an output that stage_patch will not
                    # publish. Leave even a pre-existing destination untouched.
                    result.update(success=False, error="No-op comment resolution cannot publish a distinct output_path")
                changed = Path(staged).read_bytes() != before
                accepted = not result.get("error") and result.get("success", True)
                result.setdefault("success", bool(accepted))
                # One operation per generic tool call; preserve its own detailed fields.
                result["results"] = [{"target": values.get("target") or values.get("table_id") or method.__name__[5:],
                                      "success": bool(accepted and changed)}]
                return result

            return stage_patch(source, output, mode, 1, apply)

        return wrapped
    return decorate(function) if function else decorate


def _public_paths(value: Any, staged: str, source: str) -> Any:
    if isinstance(value, str):
        return value.replace(staged, source)
    if isinstance(value, list):
        return [_public_paths(item, staged, source) for item in value]
    if isinstance(value, dict):
        return {key: _public_paths(item, staged, source) for key, item in value.items()}
    return value


_lock_guard = threading.Lock()
_path_locks: dict[str, list[Any]] = {}


@contextmanager
def _writer_locks(paths):
    """Serialise this process's writers; do not claim an external-editor lock."""
    keys = sorted({str(path) for path in paths})
    with _lock_guard:
        entries = []
        for key in keys:
            entry = _path_locks.setdefault(key, [threading.RLock(), 0])
            entry[1] += 1
            entries.append((key, entry))
    acquired = []
    try:
        for _, entry in entries:
            while not entry[0].acquire(timeout=0.1):
                raise_if_cancelled()
            acquired.append(entry[0])
            raise_if_cancelled()
        yield
    finally:
        for lock in reversed(acquired):
            lock.release()
        with _lock_guard:
            for key, entry in entries:
                entry[1] -= 1
                if not entry[1]:
                    del _path_locks[key]


def fingerprint(path: Path):
    if not path.exists():
        return None
    before = path.stat()
    if not path.is_file():
        raise ValueError("Document path must be a regular file")
    if before.st_nlink > 1:
        raise ValueError("Hard-linked document paths are unsupported for mutation")
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            raise_if_cancelled()
            digest.update(chunk)
    after = path.stat()
    def state(stat):
        return (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)

    if state(before) != state(after):
        raise ValueError("Document changed while being fingerprinted")
    return (digest.hexdigest(), state(after))


def stage_patch(
    source: str,
    output: str | None,
    mode: str,
    requested: int,
    apply: Callable[[str], dict[str, Any]],
) -> dict[str, Any]:
    """Run the existing format writers on a private copy and publish at most once.

    Preview executes the same transformations/validation as a real write, then discards
    the copy. The caller never sees a partially written source or existing destination.
    """
    source_path = Path(source).resolve()
    destination = Path(output or source).resolve()
    with _writer_locks((source_path, destination)):
        return _stage_patch_locked(source, output, mode, requested, apply, source_path, destination)


def _stage_patch_locked(source, output, mode, requested, apply, source_path, destination):
    result: dict[str, Any] = {}
    planned = 0
    committed = 0
    source_state = None
    try:
        raise_if_cancelled()
        source_state = fingerprint(source_path)
        if source_state is None:
            raise FileNotFoundError(f"File not found: {source}")
        destination_state = fingerprint(destination)
        from .package_guard import admit_package

        admit_package(source_path)
        # Same filesystem for atomic publication; the whole private directory is removed
        # on every return, including scratch files made by legacy format writers.
        with tempfile.TemporaryDirectory(prefix=".office-patch-", dir=destination.parent) as tmp:
            staged = Path(tmp) / ("document" + source_path.suffix)
            shutil.copy2(source_path, staged)
            if fingerprint(staged)[0] != source_state[0] or fingerprint(source_path) != source_state:
                raise ValueError("Source changed while staging")
            token = _staged_paths.set(_staged_paths.get() | {str(staged.resolve())})
            try:
                result = _public_paths(apply(str(staged)), str(staged), source)
            finally:
                _staged_paths.reset(token)
            raise_if_cancelled()
            planned = sum(bool(item.get("success")) for item in result.get("results", []))
            rejected = bool(result.get("errors") or result.get("error") or result.get("skipped_targets") or result.get("unmatched_targets"))
            if mode == "strict" and (rejected or planned != requested or not result.get("success")):
                result.update(success=False, status="failed")
                result.setdefault("warnings", []).append("Strict batch refused; no changes were committed.")
            elif result.get("error"):
                result.update(success=False, status="failed")
            elif planned:
                raise_if_cancelled()
                if fingerprint(source_path) != source_state:
                    raise ValueError("Source changed before commit; retry with fresh inspection")
                if staged.suffix.lower() in {".pptx", ".docx"}:
                    from .package_preservation import restore_unchanged_parts

                    restore_unchanged_parts(source_path, staged)
                raise_if_cancelled()
                validate_staged_document(staged)
                raise_if_cancelled()
                from .package_preservation import diff_package

                result["package_diff"] = diff_package(source_path, staged)
                if (Path(source).resolve() != source_path or Path(output or source).resolve() != destination
                        or fingerprint(source_path) != source_state
                        or fingerprint(destination) != destination_state):
                    raise ValueError("Source or destination changed before commit; retry with fresh inspection")
                if mode != "dry_run":
                    raise_if_cancelled()
                    os.replace(staged, destination)
                    committed = planned
            if mode == "dry_run":
                result.setdefault("warnings", []).append("Preview only; no changes were committed.")
    except MCPRequestCancelled:
        raise
    except Exception as exc:
        result.update(success=False, status="failed", error=f"Patch not committed: {exc}")
    result.update(
        mode=mode,
        file=str(output or source),
        changes_planned=planned,
        changes_applied=committed,
        source_sha256=source_state[0] if source_state else None,
    )
    result.setdefault("diagnostics", {}).update(changes_planned=planned, changes_applied=committed)
    # Retain skipped_targets for old clients while exposing every missing target explicitly.
    unmatched = result.setdefault("unmatched_targets", [])
    known = {item.get("target") for item in unmatched}
    for item in result.get("skipped_targets", []):
        if item.get("target") not in known:
            unmatched.append({"target": item.get("target"), "reason": "target_not_matched"})
    for item in result.get("results", []):
        item["applied"] = bool(committed and item.get("success"))
    return result
